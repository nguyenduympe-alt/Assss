"""Đọc thời khoá biểu / phân phối chương trình từ Word, Excel hoặc PDF — chạy NỘI BỘ.

Không gửi tệp của giáo viên ra ngoài. Không phải mô hình ngôn ngữ lớn (VPS 1 lõi / ~1 GB
RAM không chạy được LLM/vision). Bộ này:

  1. Lấy chữ từ Word (.docx), Excel (.xlsx) hoặc PDF có lớp chữ (nội bộ, không gửi ra ngoài).
  2. Phân tích theo mẫu thời khoá biểu / PPCT phổ thông Việt Nam.
  3. Trả về danh sách tiết (lớp · môn · khối) hoặc dòng bài PPCT để giáo viên duyệt rồi lưu.

Thầy/cô vẫn phải xem lại trước khi ghi — chữ trên ảnh mờ có thể đọc sai.
"""
from __future__ import annotations

import csv
import io
import os
import re
import shutil
import subprocess
import tempfile

from . import anh_ocr as AO
from . import mon_day as MD

TOI_DA = 8 * 1024 * 1024
TOI_DA_TIET = 80
TOI_DA_PPCT = 400
# Sheet "chỉ để xem" trong file mẫu — không đọc dữ liệu (Ví dụ, Hướng dẫn...)
_BO_SHEET = re.compile(r"vi\s*du|example|khong\s*doc|huong\s*dan|guide|instructions")


def _sheet_bo_qua(ten_s):
    return bool(_BO_SHEET.search(_bo_dau(ten_s or "")))
DUOI_ANH = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")
DUOI_PDF = (".pdf",)
DUOI_WORD = (".docx",)
DUOI_EXCEL = (".xlsx", ".xlsm", ".xls", ".csv")
DUOI_OK = DUOI_PDF + DUOI_ANH          # PPCT: PDF/ảnh
DUOI_TKB = DUOI_PDF + DUOI_WORD + DUOI_EXCEL

# Môn hay gặp thêm (ngoài GOI_Y) trên TKB giáo viên chủ nhiệm
MON_THEM = ("Sinh hoạt lớp", "Chào cờ", "Sinh hoạt", "SHL", "GDTC", "HĐTN",
            "Hoạt động trải nghiệm", "Tư vấn", "Thư viện", "Công dân")


def _gon(s):
    if not isinstance(s, str):
        s = "" if s is None else str(s)
    return re.sub(r"\s+", " ", s.replace("\u00a0", " ")).strip()


def _bo_dau(s):
    return MD.bo_dau(_gon(s))


def la_anh(ten):
    return (ten or "").lower().endswith(DUOI_ANH)


def la_pdf(ten):
    return (ten or "").lower().endswith(DUOI_PDF)


def la_word(ten):
    return (ten or "").lower().endswith(DUOI_WORD)


def la_excel(ten):
    return (ten or "").lower().endswith(DUOI_EXCEL)


def la_tkb(ten):
    return (ten or "").lower().endswith(DUOI_TKB)


# ---------------------------------------------------------------- lấy chữ
def _pdf_pypdf(data):
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        rd = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in rd.pages[:30])
    except Exception:
        return ""


def _pdf_stream(data):
    """Lấy chuỗi (...) Tj trong luồng PDF — đủ cho PDF do máy tạo (ReportLab, Word in PDF)."""
    import zlib
    ra = []
    for m in re.finditer(rb"stream\r?\n(.{1,800000}?)\r?\nendstream", data, re.S):
        raw = m.group(1)
        try:
            raw = zlib.decompress(raw)
        except Exception:
            pass
        try:
            s = raw.decode("latin-1", "replace")
        except Exception:
            continue
        for tm in re.finditer(r"\((?:\\.|[^\\)])*\)\s*Tj", s):
            chu = tm.group(0)[:-2].strip()
            if chu.startswith("(") and chu.endswith(")"):
                chu = chu[1:-1]
            chu = (chu.replace(r"\n", " ").replace(r"\r", " ").replace(r"\t", " ")
                      .replace(r"\(", "(").replace(r"\)", ")").replace(r"\\", "\\"))
            if chu.strip():
                ra.append(chu)
        for tm in re.finditer(r"\[(.*?)\]\s*TJ", s, re.S):
            for om in re.finditer(r"\((?:\\.|[^\\)])*\)", tm.group(1)):
                chu = om.group(0)[1:-1]
                chu = chu.replace(r"\n", " ").replace(r"\(", "(").replace(r"\)", ")")
                if chu.strip():
                    ra.append(chu)
    return "\n".join(ra)


def _ocr_tesseract(data, ten=""):
    """OCR ảnh bằng tesseract CLI nếu có — không gửi ra ngoài."""
    if not shutil.which("tesseract"):
        return ""
    duoi = ".png"
    low = (ten or "").lower()
    for d in DUOI_ANH:
        if low.endswith(d):
            duoi = d
            break
    with tempfile.NamedTemporaryFile(suffix=duoi, delete=True) as f:
        f.write(data)
        f.flush()
        try:
            p = subprocess.run(
                ["tesseract", f.name, "stdout", "-l", "vie+eng", "--psm", "6"],
                capture_output=True, timeout=45, check=False)
        except Exception:
            return ""
        out = (p.stdout or b"").decode("utf-8", "replace")
        return out



def _dong_o(ds):
    o = [_gon("" if v is None else str(v)) for v in ds]
    o = ["" if x.lower() == "nan" else x for x in o]
    return "  ".join(o) if any(o) else ""


def _docx_chu(data):
    """Chữ Word: đoạn văn + bảng (ô cách nhau, giữ lưới TKB)."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
    except Exception:
        return ""
    ra = []
    for p in doc.paragraphs:
        x = _gon(p.text)
        if x:
            ra.append(x)
    for tbl in doc.tables:
        for row in tbl.rows:
            d = _dong_o(c.text for c in row.cells)
            if d:
                ra.append(d)
    return "\n".join(ra)


def _excel_chu(data, ten=""):
    """Chữ Excel/CSV: mọi ô, không lấy hàng đầu làm tiêu đề (lưới TKB)."""
    try:
        import pandas as pd
    except ImportError:
        return ""
    bio = io.BytesIO(data)
    low = (ten or "").lower()
    try:
        if low.endswith(".csv"):
            sheets = {"_": pd.read_csv(bio, header=None, dtype=str, encoding="utf-8",
                                       engine="python")}
        else:
            sheets = pd.read_excel(bio, sheet_name=None, header=None, dtype=str)
    except Exception:
        return ""
    ra = []
    for ten_s, df in (sheets or {}).items():
        if df is None or getattr(df, "empty", True):
            continue
        if _sheet_bo_qua(ten_s):
            continue  # sheet ví dụ/hướng dẫn trong file mẫu — không đọc
        for _, row in df.iterrows():
            d = _dong_o(row.tolist())
            if d:
                ra.append(d)
    return "\n".join(ra)


def _giu_dong(s):
    """Gọn khoảng trắng nhưng GIỮ xuống dòng — parser TKB/PPCT tách theo dòng."""
    ra = []
    for x in (s or "").replace("\u00a0", " ").splitlines():
        x = re.sub(r"[ \t]+", " ", x).strip()
        if x:
            ra.append(x)
    return "\n".join(ra)


def _chen_dong(chu):
    """PDF đôi khi dồn nhiều tiết một dòng — tách trước Thứ / Tuần."""
    chu = chu or ""
    chu = re.sub(r"\s+(?=(?:thứ|thu)\s+(?:hai|ba|t[uư]|năm|nam|sáu|sau|bảy|bay|[2-8])\b)",
                 "\n", chu, flags=re.I)
    chu = re.sub(r"\s+(?=(?:tuần|tuan)\s+\d+)", "\n", chu, flags=re.I)
    return chu


def doc_chu(data, ten="", ocr=True):
    """Trả về (text, nguon, loi). nguon: pdf / word / excel / ocr / rong.

    ocr=False: không đọc ảnh / PDF scan (dùng cho thời khoá biểu).
    """
    ten = ten or ""
    if not data:
        return "", "", "Tệp trống."
    if len(data) > TOI_DA:
        return "", "", "Tệp lớn hơn 8 MB — thầy/cô lưu gọn lại rồi thử lại."
    if (ten or "").lower().endswith(".doc") and not la_word(ten):
        return "", "", "Không đọc tệp .doc cũ — thầy/cô lưu lại .docx rồi tải lên."
    if la_pdf(ten) or (data[:5] == b"%PDF-"):
        chu = _pdf_pypdf(data) or _pdf_stream(data)
        chu = _chen_dong(_giu_dong(chu))
        if len(re.sub(r"\s+", "", chu)) >= 8:
            return chu, "pdf", ""
        if ocr:
            ocr_chu = _chen_dong(_giu_dong(_ocr_tesseract(data, ten=".pdf")))
            if len(re.sub(r"\s+", "", ocr_chu or "")) >= 8:
                return ocr_chu, "ocr", ""
        return chu, "pdf", ("PDF này gần như không có chữ để đọc (thường là bản scan). "
                            "Thầy/cô xuất lại PDF từ Word, hoặc gửi tệp .docx / .xlsx.")
    if la_excel(ten):
        chu = _giu_dong(_excel_chu(data, ten))
        if len(re.sub(r"\s+", "", chu or "")) >= 8:
            return chu, "excel", ""
        return "", "excel", ("Không đọc được bảng tính. Hãy lưu .xlsx hoặc .csv, "
                             "các cột Thứ / Tiết / lớp · môn.")
    if la_word(ten) or (data[:2] == b"PK" and b"word/" in data[:12000]):
        chu = _giu_dong(_docx_chu(data))
        if len(re.sub(r"\s+", "", chu or "")) >= 8:
            return chu, "word", ""
        return "", "word", ("Không đọc được chữ trong tệp Word. "
                            "Hãy lưu lại .docx, bảng TKB có cột Thứ 2…6.")
    if la_anh(ten) or data[:8] == b"\x89PNG\r\n\x1a\n" or data[:2] == b"\xff\xd8":
        if not ocr:
            return "", "", ("Không đọc ảnh. Thầy/cô gửi Word (.docx), Excel (.xlsx) "
                            "hoặc PDF có chữ.")
        chu = _chen_dong(_giu_dong(_ocr_tesseract(data, ten)))
        if len(re.sub(r"\s+", "", chu or "")) >= 8:
            return chu, "ocr", ""
        if not shutil.which("tesseract"):
            return "", "", ("Ảnh cần bộ đọc chữ Tesseract trên máy chủ (chưa cài). "
                            "Thầy/cô gửi PDF có chữ hoặc tệp Word/Excel sẽ đọc chắc hơn.")
        return "", "", "Không đọc được chữ trên ảnh — hãy chụp rõ, thẳng, đủ sáng rồi thử lại."
    return "", "", "Chỉ nhận Word (.docx), Excel (.xlsx) hoặc PDF có chữ."


# ---------------------------------------------------------------- nhận diện môn / lớp / thứ
def _ds_mon(them=()):
    ra = []
    seen = set()
    for x in list(them or []) + list(MD.GOI_Y) + list(MON_THEM):
        t = _gon(x)
        k = _bo_dau(t)
        if t and k not in seen:
            seen.add(k)
            ra.append(t)
    ra.sort(key=lambda s: -len(s))
    return ra


def _tim_mon(chu, ds_mon=()):
    s = _bo_dau(chu)
    if not s:
        return ""
    for m in _ds_mon(ds_mon):
        k = _bo_dau(m)
        if k and re.search(r"(?<![a-z0-9])" + re.escape(k) + r"(?![a-z0-9])", s):
            return MD.chuan_mon(m)
    return ""


def _lop_hop_le(lop):
    """4B2, 3A, 10A1 — không nhận 7h / 14h (giờ)."""
    lop = re.sub(r"\s+", "", lop or "")
    m = re.match(r"^([0-9]{1,2})([A-Za-z])([A-Za-z0-9])?$", lop)
    if not m:
        return ""
    n = int(m.group(1))
    if not (1 <= n <= 12):
        return ""
    if m.group(2).lower() == "h":
        return ""
    return (m.group(1) + m.group(2) + (m.group(3) or "")).upper()


def _tim_lop(chu):
    s = _gon(chu)
    m = re.search(r"(?:lop|lớp)\s*([0-9]{1,2}\s*[A-Za-z][A-Za-z0-9]?)", s, re.I)
    if m:
        return _lop_hop_le(m.group(1))
    for m in re.finditer(r"\b([0-9]{1,2}[A-Za-z][A-Za-z0-9]?)\b", s):
        lop = _lop_hop_le(m.group(1))
        if lop:
            return lop
    return ""


def _tim_phong(chu):
    s = _gon(chu)
    m = re.search(r"((?:P\.|phòng|phong)\s*[A-Za-zÀ-ỹ0-9]+(?:\s+[A-D])?)", s, re.I)
    if m:
        return _gon(m.group(1))[:120]
    return ""


def _la_diem_khac(chu):
    """Ô 'Mỹ Phước D' / MPD — điểm trường khác, không phải lớp."""
    if _tim_lop(chu):
        return False
    s = _bo_dau(chu)
    return bool(re.search(r"my phuoc|\bmpd\b|\bmpe\b|diem truong|co so 2", s))


def _bo_gio(s):
    s = s or ""
    s = re.sub(r"\d{1,2}\s*h(?:\s*\d{1,2})?\s*(?:→|->|–|-|~)\s*\d{1,2}\s*h(?:\s*\d{1,2})?",
               " ", s, flags=re.I)
    s = re.sub(r"\d{1,2}\s*h(?:\s*\d{1,2})?", " ", s, flags=re.I)
    return s


def _o_hoc(chu, ds_mon=()):
    """Một ô TKB: 'Tin học 4B2 — P.Trường' → lớp, môn, phòng, khối."""
    chu = _gon(chu)
    if not chu or _la_diem_khac(chu):
        return None
    lop = _tim_lop(chu)
    mon = _tim_mon(chu, ds_mon)
    if not (lop and mon):
        return None
    return {"lop": lop, "mon": mon, "phong": _tim_phong(chu), "khoi": _khoi_tu_lop(lop, chu)}


def _khoi_tu_lop(lop, chu=""):
    m = re.search(r"(?:khoi|khối|lop|lớp)\s*([0-9]{1,2})\b", _gon(chu), re.I)
    if m and 1 <= int(m.group(1)) <= 12:
        return str(int(m.group(1)))
    m = re.match(r"([0-9]{1,2})", lop or "")
    if m and 1 <= int(m.group(1)) <= 12:
        return str(int(m.group(1)))
    return ""


def _tim_thu(chu):
    s = _bo_dau(chu)
    bang = (("chu nhat", 8), ("cn", 8), ("thu bay", 7), ("thu 7", 7), ("t7", 7),
            ("thu sau", 6), ("thu 6", 6), ("t6", 6),
            ("thu nam", 5), ("thu 5", 5), ("t5", 5),
            ("thu tu", 4), ("thu 4", 4), ("t4", 4),
            ("thu ba", 3), ("thu 3", 3), ("t3", 3),
            ("thu hai", 2), ("thu 2", 2), ("t2", 2))
    for k, v in bang:
        if re.search(r"(?<![a-z0-9])" + re.escape(k) + r"(?![a-z0-9])", s):
            return v
    return None


def _tim_tiet(chu):
    s = _bo_dau(chu)
    m = re.search(r"(?:tiet|tiết)\s*([0-9]{1,2})", s)
    if m and 1 <= int(m.group(1)) <= 10:
        return int(m.group(1))
    return None


def _tim_buoi(chu):
    s = _bo_dau(chu)
    if re.search(r"(?<![a-z0-9])chieu(?![a-z0-9])", s):
        return "Chiều"
    if re.search(r"(?<![a-z0-9])sang(?![a-z0-9])", s):
        return "Sáng"
    return ""


# ---------------------------------------------------------------- TKB
def _dong_tkb(thu, buoi, tiet, lop, mon, khoi, phong=""):
    if not (thu and tiet and lop and mon):
        return None
    khoi = khoi or _khoi_tu_lop(lop)
    if khoi not in {str(i) for i in range(1, 13)}:
        return None
    return {"thu": int(thu), "buoi": buoi or "Sáng", "tiet": int(tiet),
            "lop": lop[:60], "mon": mon[:100], "khoi": khoi, "phong": (phong or "")[:120]}


def phan_tich_tkb(chu, ds_mon=()):
    """Phân tích chữ → danh sách tiết dạy (tối đa TOI_DA_TIET)."""
    chu = chu or ""
    ra, seen = [], set()

    def them(d):
        if not d:
            return
        k = (d["thu"], d["buoi"], d["tiet"], d["lop"], _bo_dau(d["mon"]))
        if k in seen or len(ra) >= TOI_DA_TIET:
            return
        seen.add(k)
        ra.append(d)

    # 1) dòng đủ trường: "Thứ 2, Sáng, Tiết 1, lớp 6A, Toán"
    for dong in re.split(r"[\n;]+", chu):
        dong = _gon(dong)
        if len(dong) < 6:
            continue
        thu = _tim_thu(dong)
        tiet = _tim_tiet(dong)
        o = _o_hoc(dong, ds_mon)
        if not (thu and tiet and o):
            continue
        buoi = _tim_buoi(dong) or "Sáng"
        them(_dong_tkb(thu, buoi, tiet, o["lop"], o["mon"], o["khoi"], o["phong"]))

    # 2) lưới Thứ × Tiết (ảnh/PDF kiểu: Tin học 4B2 — P.Trường)
    _phan_tich_luoi(chu, ds_mon, them)

    ra.sort(key=lambda x: (x["thu"], 0 if x["buoi"] == "Sáng" else 1, x["tiet"], x["lop"]))
    return ra


def _thus_header(dong):
    """[(thu, vị_trí), ...] từ dòng 'Thứ 2  Thứ 3  Thứ 4 …'."""
    ra = []
    for m in re.finditer(
            r"(?:thứ|thu)\s*(?:hai|ba|tư|tu|năm|nam|sáu|sau|bảy|bay|[2-8])",
            dong, re.I):
        t = _tim_thu(m.group(0))
        if t:
            ra.append((t, m.start()))
    return ra


def _gop_dong_o(lines):
    """Gộp dòng bọc ô: 'Tin học' + '4B2 — P.Trường' thành một ô."""
    ra = []
    for x in lines:
        g = _gon(x)
        if not g:
            continue
        if ra and (re.match(r"^\d{1,2}[A-Za-z]", g)
                   or re.match(r"^(?:P\.|phòng|phong|—|-|–)", g, re.I)):
            ra[-1] = ra[-1].rstrip() + " " + g
        else:
            ra.append(x.rstrip())
    return ra


def _tiet_buoi_dong(dong, buoi_mac="Sáng"):
    s = _bo_gio(dong)
    buoi = _tim_buoi(s) or buoi_mac
    s2 = re.sub(r"(?i)(?:sáng|sang|chiều|chieu|buổi|buoi|tiết|tiet)", " ", s)
    m = re.search(r"(?:^|\s)(\d{1,2})\b", s2.strip())
    tiet = None
    if m and 1 <= int(m.group(1)) <= 10:
        tiet = int(m.group(1))
    t2 = _tim_tiet(dong)
    if t2:
        tiet = t2
    return tiet, buoi


def _o_trong_dong(dong, ds_mon=()):
    """Các ô có môn + lớp trên một dòng, kèm vị trí — không lấy nhầm lớp ô bên cạnh."""
    ra = []
    for m in re.finditer(r"\b(\d{1,2}[A-Za-z][A-Za-z0-9]?)\b", dong):
        lop = _lop_hop_le(m.group(1))
        if not lop:
            continue
        truoc = re.split(r"\b\d{1,2}[A-Za-z][A-Za-z0-9]?\b", dong[:m.start()])[-1][-40:]
        sau = re.split(r"\b\d{1,2}[A-Za-z]", dong[m.end():], 1)[0][:36]
        # 'Tin học 4B2' (môn trước) hoặc '6A Toan' (môn sau)
        mon = _tim_mon(sau, ds_mon) or _tim_mon(truoc, ds_mon)
        if not mon:
            continue
        ra.append({"lop": lop, "mon": mon, "phong": _tim_phong(sau),
                   "khoi": _khoi_tu_lop(lop), "start": m.start(), "raw": (truoc + lop + sau)[-80:]})
    return ra


def _map_thu_vitri(thus_pos, start):
    if not thus_pos:
        return None
    chon, _ = thus_pos[0]
    best = 10 ** 9
    for thu, st in thus_pos:
        d = abs(start - st)
        if start >= st - 3:
            chon = thu
        if d < best:
            best, chon_n = d, thu
    # ưu tiên cột có mốc nằm trái ô; nếu lệch quá xa lấy cột gần nhất
    if best <= 18:
        return chon_n
    return chon


def _map_thu_neo(thus, tokens):
    """Cột điểm trường (Mỹ Phước D) làm mốc — thường là cả cột Thứ 5."""
    if not thus or not tokens:
        return []
    neo = [i for i, tk in enumerate(tokens) if _la_diem_khac(tk.get("raw") or tk.get("lop") or "")]
    # tokens here are o_hoc dicts; campus rows stored as {"raw": ..., "campus": True}
    neo = [i for i, tk in enumerate(tokens) if tk.get("campus")]
    if len(neo) != 1:
        return []
    thu_neo = 5 if 5 in thus else thus[min(len(thus) - 1, max(0, len(thus) // 2 + 1))]
    i_neo, i_thu = neo[0], thus.index(thu_neo)
    ra = []
    for i, tk in enumerate(tokens):
        j = i_thu - i_neo + i
        if 0 <= j < len(thus) and not tk.get("campus"):
            ra.append((thus[j], tk))
    return ra


def _token_dong(dong, ds_mon):
    """Ô học + ô điểm trường, theo thứ tự trái → phải."""
    moc = []
    for m in re.finditer(r"mỹ\s*phước\s*\w+|\bmpd\b|\bmpe\b", dong, re.I):
        moc.append({"campus": True, "start": m.start(), "raw": m.group(0)})
    moc.extend(_o_trong_dong(dong, ds_mon))
    moc.sort(key=lambda x: x.get("start", 0))
    return moc


def _phan_tich_luoi(chu, ds_mon, them):
    raw = _gop_dong_o([x for x in (chu or "").splitlines() if x.strip()])
    header_i, thus_pos = -1, []
    for i, dong in enumerate(raw):
        tp = _thus_header(dong)
        if len(tp) >= 3:
            header_i, thus_pos = i, tp
            break
    if header_i < 0:
        return
    thus = [t for t, _ in thus_pos]
    buoi = "Sáng"
    for dong in raw[header_i + 1: header_i + 24]:
        tiet, buoi2 = _tiet_buoi_dong(dong, buoi)
        if _tim_buoi(dong):
            buoi = buoi2
        if not tiet:
            if _tim_buoi(dong) == "Chiều":
                buoi = "Chiều"
            continue
        buoi = buoi2 or buoi
        tokens = _token_dong(_bo_gio(dong), ds_mon)
        o_hoc = [tk for tk in tokens if not tk.get("campus")]
        gan = _map_thu_neo(thus, tokens)
        if not gan and len(o_hoc) == len(thus):
            gan = list(zip(thus, o_hoc))
        if not gan and o_hoc and thus_pos and len(dong) >= max(40, (thus_pos[-1][1] or 0) * 0.5):
            for tk in o_hoc:
                thu = _map_thu_vitri(thus_pos, tk.get("start") or 0)
                if thu:
                    gan.append((thu, tk))
        for thu, tk in gan:
            if tk.get("campus"):
                continue
            them(_dong_tkb(thu, buoi, tiet, tk.get("lop"), tk.get("mon"),
                           tk.get("khoi"), tk.get("phong")))


# ---------------- TKB bảng tính dạng LƯỚI (Buổi | Tiết | Thứ 2..7) ----------------
def _luoi_bang_tinh(data, ten=""):
    """Mỗi sheet → (tên_sheet, lưới ô chuỗi). Hỗ trợ .xlsx/.xlsm, .xls, .csv."""
    ten = (ten or "").lower()
    if ten.endswith(".csv"):
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except Exception:
                continue
        if text is None:
            text = data.decode("utf-8", "replace")
        try:
            nl = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except Exception:
            nl = csv.excel
        return [("_", [list(hang) for hang in csv.reader(io.StringIO(text), nl)])]
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        try:
            import pandas as pd
        except ImportError:
            return []
        ra = []
        try:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=str)
            for ten_s, df in (sheets or {}).items():
                ra.append((ten_s, [[_gon(v) for v in hang]
                                   for hang in df.fillna("").values.tolist()]))
        except Exception:
            return []
        return [(t, g) for t, g in ra if g]
    ra = []
    try:
        for ws in wb.worksheets:
            g = [[_gon(v) for v in hang]
                 for hang in ws.iter_rows(values_only=True, max_row=140, max_col=40)]
            while g and not any(any(x for x in h) for h in g[-1]):
                g.pop()
            if g:
                ra.append((ws.title, g))
    finally:
        wb.close()
    return ra


def _tim_hang_thu(g):
    """Hàng tiêu đề + {cột: thứ} khi một hàng có ≥3 ô 'Thứ …' ngắn."""
    for i, dong in enumerate(g[:40]):
        thu_cot = {}
        for j, o in enumerate(dong):
            o_g = _gon(o)
            if not o_g or len(o_g) > 14:
                continue
            t = _tim_thu(o_g)
            if t:
                thu_cot[j] = t
        if len(thu_cot) >= 3:
            return i, thu_cot
    return -1, {}


def _doc_luoi_grid(g, hi, thu_cot, ds_mon, seen, ra):
    """Vòng lặp dùng chung: từ hàng tiêu đề hi + map cột→thứ, đọc từng ô học."""
    cot_buoi = cot_tiet = None
    for j, o in enumerate(g[hi]):
        n = _bo_dau(o)
        if cot_buoi is None and re.search(r"buoi", n):
            cot_buoi = j
        if cot_tiet is None and re.search(r"^tiet\b", n):
            cot_tiet = j
    buoi = "Sáng"
    for dong in g[hi + 1:]:
        if cot_buoi is not None and cot_buoi < len(dong):
            b = _tim_buoi(dong[cot_buoi])
            if b:
                buoi = b
        tiet = None
        for j in ((cot_tiet,) if cot_tiet is not None else (1, 0, 2)):
            if j is None or j >= len(dong):
                continue
            o = re.sub(r"\.0$", "", _gon(dong[j]))
            m = re.match(r"^(?:tiet\s*)?([0-9]{1,2})$", _bo_dau(o), re.I)
            if m and 1 <= int(m.group(1)) <= 10:
                tiet = int(m.group(1))
                break
        if not tiet:
            continue
        for j, thu in thu_cot.items():
            o = dong[j] if j < len(dong) else ""
            if not o:
                continue
            tk = _o_hoc(o, ds_mon)
            if not tk:
                continue
            r = _dong_tkb(thu, buoi, tiet, tk.get("lop"), tk.get("mon"),
                          tk.get("khoi"), tk.get("phong"))
            if r:
                k = (r["thu"], r["buoi"], r["tiet"], r["lop"], r["mon"])
                if k not in seen:
                    seen.add(k)
                    ra.append(r)
        if len(ra) >= TOI_DA_TIET:
            break
    return ra


def phan_tich_luoi_oo(data, ds_mon=()):
    """Đọc TKB Excel THEO Ô LƯỚI: cột Buổi, Tiết, Thứ 2..7 — mỗi ô học là 1 tiết.

    Trả danh sách tiết đúng cột Thứ (không lệch cột như cách phân tích chữ).
    Ô 'Mỹ Phước D' (điểm trường), ô trống và sheet 'Vi du' (chỉ để tham khảo) tự bỏ qua."""
    ra, seen = [], set()
    for ten_sheet, g in _luoi_bang_tinh(data):
        if _sheet_bo_qua(ten_sheet):
            continue  # sheet ví dụ/hướng dẫn trong file mẫu — không đọc
        hi, thu_cot = _tim_hang_thu(g)
        if hi < 0:
            continue
        _doc_luoi_grid(g, hi, thu_cot, ds_mon, seen, ra)
        if len(ra) >= TOI_DA_TIET:
            break
    return ra


def phan_tich_luoi_docx(data, ds_mon=()):
    """Đọc TKB Word là BẢNG lưới (Buổi | Tiết | Thứ 2..7) — theo ô, đúng cột."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
    except Exception:
        return []
    ra, seen = [], set()
    for tbl in doc.tables:
        g = [[_gon(c.text) for c in row.cells] for row in tbl.rows]
        hi, thu_cot = _tim_hang_thu(g)
        if hi < 0:
            continue
        _doc_luoi_grid(g, hi, thu_cot, ds_mon, seen, ra)
        if len(ra) >= TOI_DA_TIET:
            break
    return ra


def _lop_uan(lop):
    """Lớp có chữ cái hợp lý — loại O/I/Y/J/Q/U (OCR rác hay sinh 3OY, 5Q1...)."""
    m = re.match(r"^([0-9]{1,2})([A-Za-z])([A-Za-z0-9])?$", lop or "")
    return bool(m) and m.group(2).upper() not in "OIYJQU"


def phan_tich_anh_tkb(lines, ds_mon=()):
    """Dựng tiết dạy từ OCR tokens CÓ TỌA ĐỘ của ảnh TKB lưới Buổi × Tiết × Thứ 2..7.

    Gán từng dòng chữ: cột Thứ theo MÉP TRÁI (x0 ≈ mép cột), tiết theo mốc cột số
    (so ĐỈNH dòng y0), buổi theo ranh giới Sáng/Chiều — rồi gom các dòng cùng ô.
    Ô 'Mỹ Phước D' (điểm trường), tiêu đề, cột số tự loại."""
    lines = [dict(l) for l in (lines or []) if _gon(l.get("t"))]
    if not lines:
        return []
    for l in lines:
        l["cx"] = (l["x0"] + l["x1"]) / 2.0
        l["cy"] = (l["y0"] + l["y1"]) / 2.0
        l["h"] = max(1.0, l["y1"] - l["y0"])
    # 1) cột Thứ từ tiêu đề "Thứ N"; mốc cột Tiết
    thu_cot, cot_tiet_x = [], None
    for l in lines:
        s = _bo_dau(l["t"]).replace(" ", "")
        # OCR hay đọc lệch: "Thứ 3"→"Thuú 3", "Thứ 4"→"Thur 4" — nới khớp th+chữ+[2-8]
        m = re.match(r"^th[a-z]{0,3}([2-8])$", s)
        if m and len(s) <= 7:
            thu = int(m.group(1))
            if all(thu != t for t, _ in thu_cot):
                thu_cot.append((thu, l["cx"]))
        if cot_tiet_x is None and s == "tiet":
            cot_tiet_x = l["cx"]
    if not thu_cot:
        return []
    thu_cot.sort(key=lambda z: z[1])
    # OCR mất mốc một cột giữa → nội suy đều các cột thiếu từ khe hở
    day_du = []
    for (ta, xa), (tb, xb) in zip(thu_cot, thu_cot[1:]):
        day_du.append((ta, xa))
        for k in range(1, tb - ta):
            day_du.append((ta + k, xa + (xb - xa) * k / (tb - ta)))
    day_du.append(thu_cot[-1])
    thu_cot = day_du
    xs = [c for _, c in thu_cot]
    buoc_cot = min((b - a for a, b in zip(xs, xs[1:])), default=120.0) or 120.0
    if cot_tiet_x is None:
        cot_tiet_x = min(xs) - buoc_cot

    # 2) mốc y các tiết: token thuần số (1..10) ở cột Tiết — so theo đỉnh dòng
    tiet_anchors = []
    for l in lines:
        s = _bo_dau(l["t"]).replace(" ", "")
        m = re.match(r"^(?:tiet)?([0-9]|10)$", s)
        if not m or l["h"] > buoc_cot * 0.6:
            continue
        if abs(l["cx"] - cot_tiet_x) < buoc_cot * 0.6:
            tiet_anchors.append((int(m.group(1)), l["y0"]))
    # 3) ranh giới Sáng/Chiều
    y_sang = y_chieu = None
    for l in lines:
        s = _bo_dau(l["t"]).replace(" ", "")
        if l["x0"] < min(xs):
            if s.endswith("sang"):
                y_sang = l["cy"]
            elif s.endswith("chieu"):
                y_chieu = l["cy"]

    def buoi_cua(cy):
        if y_sang is not None and y_chieu is not None:
            return "Sáng" if cy < (y_sang + y_chieu) / 2.0 else "Chiều"
        if y_chieu is not None:
            return "Sáng" if cy < y_chieu - 5 else "Chiều"
        return "Sáng"

    def tiet_cua(y0):
        if not tiet_anchors:
            return None
        t, y_a = min(tiet_anchors, key=lambda z: abs(z[1] - y0))
        if y_a - y0 > buoc_cot * 0.7:   # dòng phải nằm dưới mốc tiết của nó
            return None
        return t

    # 4) gán từng token (thu, buoi, tiet) rồi gom các dòng cùng ô
    khong_cho = set()
    for l in lines:
        s = _bo_dau(l["t"]).replace(" ", "")
        if re.match(r"^th[a-z]{0,3}([2-8])$", s) and len(s) <= 7:
            khong_cho.add(id(l))
        elif s in ("buoi", "tiet") or re.match(r"^(?:tiet)?([0-9]|10)$", s):
            khong_cho.add(id(l))
        elif l["x0"] < min(xs) - buoc_cot * 0.45:
            khong_cho.add(id(l))
    o = {}
    for l in lines:
        if id(l) in khong_cho:
            continue
        thu = min(thu_cot, key=lambda z: abs(l["x0"] - z[1]))[0]
        x_cot = dict(thu_cot)[thu]
        if abs(l["x0"] - x_cot) > buoc_cot * 0.55:
            continue
        tiet = tiet_cua(l["y0"])
        if not tiet:
            continue
        k = (thu, buoi_cua(l["cy"]), tiet)
        o.setdefault(k, []).append(l)
    ra, seen = [], set()
    con_thieu_mon = []      # ô chỉ còn lớp (dòng tên môn bị OCR sót)
    dem_mon = {}
    for (thu, buoi, tiet), g in o.items():
        g.sort(key=lambda l: (l["y0"], l["x0"]))
        chu = " ".join(_gon(l["t"]) for l in g)[:160]
        tk = _o_hoc(chu, ds_mon)
        if not tk and not _tim_lop(chu) and re.search(r"(?i)p\s*t", _bo_dau(chu)):
            # OCR hay nhầm B → 6 trong tên lớp khi ô có ghi phòng/GV: "461 - P.Thới" → "4B1"
            chu2 = re.sub(r"\b(\d{1,2})6(?=[0-9]\b)", lambda m: m.group(1) + "B", chu)
            if _tim_lop(chu2):
                tk = _o_hoc(chu2, ds_mon)
        if tk and not _lop_uan(tk.get("lop")):
            # lớp kiểu OCR rác (3OY, 5Q1...) — tìm lớp khác hợp lý hơn trong cùng ô
            for m in re.finditer(r"\b([0-9]{1,2}[A-Za-z][A-Za-z0-9]?)\b", chu):
                lop2 = _lop_hop_le(m.group(1))
                if lop2 and _lop_uan(lop2) and lop2 != tk.get("lop"):
                    tk2 = _o_hoc(chu, ds_mon)
                    tk2 = tk2 or dict(tk)
                    tk2["lop"] = lop2
                    tk = tk2
                    break
        if tk and not _lop_uan(tk.get("lop")):
            tk = None
        if tk:
            dem_mon[tk.get("mon")] = dem_mon.get(tk.get("mon"), 0) + 1
            r = _dong_tkb(thu, buoi, tiet, tk.get("lop"), tk.get("mon"),
                          tk.get("khoi"), tk.get("phong"))
            if r:
                k = (r["thu"], r["buoi"], r["tiet"], r["lop"], r["mon"])
                if k not in seen:
                    seen.add(k)
                    ra.append(r)
        else:
            lop = _tim_lop(chu)
            if lop and _lop_uan(lop):
                con_thieu_mon.append((thu, buoi, tiet, lop, _tim_phong(chu)))
    # ô thiếu tên môn → mượn môn xuất hiện nhiều nhất trong chính bảng này
    if con_thieu_mon and dem_mon:
        mon_cho = max(dem_mon, key=dem_mon.get)
        for thu, buoi, tiet, lop, phong in con_thieu_mon:
            r = _dong_tkb(thu, buoi, tiet, lop, mon_cho, "", phong)
            if r:
                k = (r["thu"], r["buoi"], r["tiet"], r["lop"], r["mon"])
                if k not in seen:
                    seen.add(k)
                    ra.append(r)
    return ra[:TOI_DA_TIET]


def mau_xlsx_tkb():
    """File mẫu Excel LƯỚI: sheet 'TKB' trống để thầy/cô điền, sheet 'Vi du' tham khảo."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Border(*(Side(style="thin", color="64748B") for _ in range(4)))
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fill_h = PatternFill("solid", fgColor="DCE9F7")     # xanh nhạt: tiêu đề
    fill_hv = PatternFill("solid", fgColor="EFF6FF")    # xanh rất nhạt: Buổi/Tiết
    fill_bt = PatternFill("solid", fgColor="D6E6F7")    # xanh: ô có tiết dạy
    font_h = Font(bold=True, size=11)
    font_do = Font(color="C00000")

    DIEM = "Mỹ Phước D"  # ô điểm trường khác — web tự bỏ qua khi đọc

    def ve_khung(ws, sang, chieu):
        """Vẽ lưới Buổi × Tiết × Thứ 2..6; dict rỗng = khung trống để điền."""
        for j, t in enumerate(["Buổi", "Tiết", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6"], 1):
            c = ws.cell(1, j, t)
            c.font, c.fill, c.alignment, c.border = font_h, fill_h, ctr, thin
        ws.row_dimensions[1].height = 20
        r = 2
        for ten_buoi, bang in (("Sáng", sang), ("Chiều", chieu)):
            dau = r
            for tiet in range(1, 6):
                ws.cell(r, 2, tiet).alignment = ctr
                cot = bang.get(tiet, {})
                for thu in range(2, 7):
                    nd = cot.get(thu, "")
                    c = ws.cell(r, thu + 1, nd)
                    if nd == DIEM:
                        c.font = font_do
                    elif nd:
                        c.fill = fill_bt
                for j in range(1, 8):
                    cc = ws.cell(r, j)
                    cc.border = thin
                    if j <= 2:
                        cc.fill = fill_hv
                ws.row_dimensions[r].height = 30
                r += 1
            ws.cell(dau, 1, ten_buoi)
            ws.merge_cells(start_row=dau, start_column=1, end_row=r - 1, end_column=1)
            cb = ws.cell(dau, 1)
            cb.font, cb.alignment = Font(bold=True), ctr
        for i, w in enumerate((10, 7, 19, 19, 19, 19, 19), 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "C2"

    def o_hoc(lop, gv):
        return "Tin học\n%s — %s" % (lop, gv)

    sang = {1: {5: DIEM},
            2: {2: o_hoc("4B2", "P.Trường"), 3: o_hoc("2B1", "P.Thới A"),
                5: DIEM, 6: o_hoc("3B1", "P.Thới A")},
            3: {5: DIEM, 6: o_hoc("5B1", "P.Thới A")},
            4: {2: o_hoc("5B3", "P.Trường"), 3: o_hoc("3B3", "P.Trường"), 5: DIEM},
            5: {}}
    chieu = {1: {2: o_hoc("5B2", "P.Thới A"), 3: o_hoc("4B1", "P.Thới A"), 5: DIEM},
             2: {5: DIEM},
             3: {2: o_hoc("2B2", "P.Trường"), 3: o_hoc("3B2", "P.Thới A"), 5: DIEM},
             4: {}, 5: {}}

    wb = Workbook()
    ws = wb.active
    ws.title = "TKB"
    ve_khung(ws, {}, {})            # KHUNG TRỐNG — thầy/cô điền lịch của mình
    vd = wb.create_sheet("Vi du")
    ve_khung(vd, sang, chieu)       # ví dụ tham khảo — web KHÔNG đọc sheet này

    hd = wb.create_sheet("Huong dan")
    hd["A1"] = (
        "FILE MẪU THỜI KHÓA BIỂU (dạng lưới)\n"
        "— Điền lịch của bạn vào sheet “TKB”: giữ nguyên hàng tiêu đề Buổi · Tiết · Thứ 2…Thứ 6.\n"
        "— Mỗi ô có tiết dạy ghi: TÊN MÔN, xuống dòng, LỚP — Phòng/GV. Ví dụ: “Tin học 4B2 — P.Trường”\n"
        "  (ghi một dòng “Tin học 4B2 — P.Trường” cũng được; lớp kiểu 4B2, 3A, 10A1).\n"
        "— Ô trống = không dạy. Ô “Mỹ Phước D” (điểm trường khác) web tự bỏ qua khi đọc.\n"
        "— Sheet “Vi du” chỉ ĐỂ XEM cách điền — web không đọc sheet này, có thể xoá đi.\n"
        "— Lưu .xlsx rồi tải lên trang Thời khoá biểu → “Đọc từ tệp” → xem trước → Lưu.\n"
        "  Cũng nhận Word/PDF lưới Thứ 2…6 và mẫu danh sách cũ (Thứ · Buổi · Tiết · Lớp · Môn · Phòng).\n"
        "  Không gửi tệp ra ngoài.")
    hd["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    hd.column_dimensions["A"].width = 95
    hd.row_dimensions[1].height = 160
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _la_mau_tkb_trong(data):
    """Là file mẫu TKB của web nhưng sheet 'TKB' chưa được điền lịch."""
    wb = None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if "TKB" not in wb.sheetnames:
            return False
        hang1 = [_gon(v) for v in next(wb["TKB"].iter_rows(values_only=True, max_row=1, max_col=7))]
        return hang1[:3] == ["Buổi", "Tiết", "Thứ 2"]
    except Exception:
        return False
    finally:
        if wb is not None:
            wb.close()


def doc_tkb(data, ten="", ds_mon=()):
    """Đọc tệp Word / PDF / Excel / ẢNH → (ok, thong_bao, rows, nguon)."""
    ten = ten or ""
    if la_anh(ten) or (not ten and data[:4] in (b"\x89PNG", b"\xff\xd8\xff\xe0",
                                                b"\xff\xd8\xff\xe1", b"\xff\xd8\xff\xdb")):
        # ẢNH CHỤP TKB — OCR tiếng Việt chạy nội bộ (RapidOCR, tiến trình riêng)
        ok, tb, lines = AO.doc_tokens(data, ten)
        if not ok:
            return False, tb, [], "Ảnh (OCR)"
        rows = phan_tich_anh_tkb(lines, ds_mon)
        if rows:
            return True, ("Đọc được %d tiết từ ảnh thời khoá biểu (OCR tiếng Việt — "
                          "kiểm tra kỹ bảng xem trước trước khi lưu)." % len(rows)), rows, "Ảnh (OCR)"
        return False, ("Đã nhận chữ trong ảnh (%s) nhưng chưa tách được tiết theo lưới "
                       "Buổi × Tiết × Thứ. Ảnh nên chụp thẳng, đủ khung có tiêu đề "
                       "Thứ 2…6, cột Buổi và Tiết." % tb), [], "Ảnh (OCR)"
    if ten and not (la_tkb(ten) or la_pdf(ten) or (data[:5] == b"%PDF-")):
        return False, "Chỉ nhận Word (.docx), Excel (.xlsx/.csv) hoặc PDF có chữ.", [], ""
    # (TKB-LUOI) bảng tính dạng lưới "Buổi | Tiết | Thứ 2..7" — đọc theo ô, đúng cột từng thứ
    if la_excel(ten):
        try:
            rows_luoi = phan_tich_luoi_oo(data, ds_mon)
        except Exception:
            rows_luoi = []
        if rows_luoi:
            return True, ("Đọc được %d tiết từ thời khoá biểu (Excel — lưới Buổi × Tiết × Thứ)."
                          % len(rows_luoi)), rows_luoi, "Excel (lưới)"
        try:
            if _la_mau_tkb_trong(data):
                return False, ("Đây là file mẫu chưa điền lịch — hãy điền tiết dạy vào sheet "
                               "“TKB” (xem sheet “Vi du”) rồi tải lên lại."), [], "Excel (lưới)"
        except Exception:
            pass
    if la_word(ten):
        try:
            rows_w = phan_tich_luoi_docx(data, ds_mon)
        except Exception:
            rows_w = []
        if rows_w:
            return True, ("Đọc được %d tiết từ thời khoá biểu (Word — lưới Buổi × Tiết × Thứ)."
                          % len(rows_w)), rows_w, "Word (lưới)"
    chu, nguon, loi = doc_chu(data, ten, ocr=False)
    if loi and not chu:
        return False, loi, [], nguon
    rows = phan_tich_tkb(chu, ds_mon)
    if not rows:
        return False, ("Đã lấy chữ nhưng chưa tách được tiết dạy. Tệp nên có dạng "
                       "“Thứ 2, Sáng, Tiết 1, lớp 6A, Toán” hoặc lưới Thứ × Tiết "
                       "(Word/Excel/PDF)."), [], nguon
    return True, "Đọc được %d tiết từ thời khoá biểu (%s)." % (len(rows), nguon), rows, nguon


# ---------------------------------------------------------------- PPCT
def phan_tich_ppct(chu):
    """Phân tích chữ → danh sách {tuan, tiet_pp, ten_bai, ghi_chu}."""
    ra, seen = [], set()
    tuan_dang = 1
    for dong in chu.splitlines():
        dong = _gon(dong)
        if len(dong) < 3:
            continue
        s = _bo_dau(dong)
        if any(x in s for x in ("tuan tiet", "ten bai", "phan phoi", "ghi chu", "ppct")):
            if "ten bai" in s or "tuan" == s.split(" ")[0] and "tiet" in s:
                continue
        m_tuan = re.search(r"(?:tuan|tuần)\s*(\d{1,2}(?:\s*[-_,;/+]\s*\d{1,2})*)", s)
        if m_tuan:
            tuan_dang = m_tuan.group(1)
        tiet = None
        m_tiet = re.search(r"(?:tiet|tiết)\s*(?:ppct\s*)?(\d{1,3}(?:\s*[-_,;/+]\s*\d{1,3})*)", s)
        if m_tiet:
            tiet = m_tiet.group(1)
        ten = dong
        ten = re.sub(r"(?i)tuần\s*\d+(?:\s*[-_,;/+]\s*\d+)*\s*[,:.\-]?\s*", "", ten)
        ten = re.sub(r"(?i)tuan\s*\d+(?:\s*[-_,;/+]\s*\d+)*\s*[,:.\-]?\s*", "", ten)
        ten = re.sub(r"(?i)tiết\s*(?:ppct\s*)?\d+(?:\s*[-_,;/+]\s*\d+)*\s*[,:.\-]?\s*", "", ten)
        ten = re.sub(r"(?i)tiet\s*(?:ppct\s*)?\d+(?:\s*[-_,;/+]\s*\d+)*\s*[,:.\-]?\s*", "", ten)
        ten = _gon(ten)
        # "1  1  Ôn tập số tự nhiên" hoặc "1-2  1+2  Tên bài"
        m3 = re.match(r"^(\d{1,2}(?:\s*[-_,;/+]\s*\d{1,2})*)\s+(\d{1,3}(?:\s*[-_,;/+]\s*\d{1,3})*)\s+(.{3,})$", dong)
        if m3:
            tuan_dang, tiet, ten = m3.group(1), m3.group(2), _gon(m3.group(3))
        if not ten or len(ten) < 3:
            continue
        if re.match(r"^\d+$", ten):
            continue
        if _bo_dau(ten) in ("tuan", "tiet", "ten bai", "ghi chu", "mon", "khoi"):
            continue
        k = (tuan_dang, tiet, _bo_dau(ten))
        if k in seen or len(ra) >= TOI_DA_PPCT:
            continue
        seen.add(k)
        ra.append({"tuan": tuan_dang, "tiet_pp": tiet, "ten_bai": ten[:300], "ghi_chu": ""})
    from . import ppct_tach as PT
    return PT.no_rong_tuan_tiet(ra)


def doc_ppct(data, ten="", mon="", khoi=""):
    """Đọc tệp → (ok, thong_bao, phan) với phan giống ppct_tach.tach (một phần)."""
    chu, nguon, loi = doc_chu(data, ten)
    if loi and not chu:
        return False, loi, []
    rows = phan_tich_ppct(chu)
    if not rows:
        return False, ("Đã lấy chữ nhưng chưa tách được bài PPCT. Tệp nên có dòng "
                       "“Tuần 1, Tiết 1, Tên bài …” hoặc “1  1  Tên bài”."), []
    ph = {"mon": mon, "khoi": khoi, "rows": rows, "loi": "", "cach": nguon}
    return True, "Đọc được %d bài PPCT (%s)." % (len(rows), nguon), [ph]
