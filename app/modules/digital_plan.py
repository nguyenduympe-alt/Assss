"""Word plans: preserve the supplied layout and provide editable teaching ideas."""
import copy
import io
import re
import unicodedata
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

from . import ai_giao_duc as AIGD

SOURCE = 'https://congbao.chinhphu.vn/van-ban/thong-tu-so-02-2025-tt-bgddt-44148.htm'
# Màu phân biệt nội dung hệ thống đề xuất: năng lực số đỏ FF0000, giáo dục AI xanh dương 0000FF
DO = RGBColor(0xFF, 0x00, 0x00)
XANH = RGBColor(0x00, 0x00, 0xFF)
MAU_COT = {'digital': DO, 'ai': XANH}


def _to_mau(cell, key, item):
    """Tô màu ô vừa ghi nếu là cột năng lực số / AI do hệ thống đề xuất (không tô nội dung gốc)."""
    mau = MAU_COT.get(key)
    if mau is None or not item.get('_themmoi', {}).get(key):
        return
    for p in cell.paragraphs:
        for run in p.runs:
            if (run.text or '').strip():
                run.font.color.rgb = mau
DOMAINS = ['Khai thác dữ liệu và thông tin', 'Giao tiếp và hợp tác trong môi trường số',
           'Sáng tạo nội dung số', 'An toàn', 'Giải quyết vấn đề', 'Ứng dụng trí tuệ nhân tạo']
TEMPLATE = Path(__file__).resolve().parents[1] / 'assets' / 'khgd-template.docx'


def plain(text):
    return ''.join(c for c in unicodedata.normalize('NFD', text.lower().replace('đ', 'd')) if not unicodedata.combining(c))


def read_word(data):
    if len(data) > 8 * 1024 * 1024:
        raise ValueError('File Word tối đa 8 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(f.file_size for f in z.infolist()) > 40 * 1024 * 1024 or len(z.infolist()) > 3000:
                raise ValueError('File Word quá lớn sau giải nén.')
            if 'word/document.xml' not in z.namelist():
                raise ValueError('File không có cấu trúc Word hợp lệ.')
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError):
        raise ValueError('Không đọc được file Word. Hãy lưu lại dưới dạng .docx.') from None
    return doc


# ---------------- (M7) Đọc tựa bài → suy luận nội dung tích hợp phù hợp ----------------
# Mỗi nhóm nội dung ứng với một mạch của khung: tựa bài khớp từ khoá nào thì hệ thống lấy
# đúng nội dung của nhóm đó (không dùng chung một câu cho mọi bài).
NHOM_TICH_HOP = [
    {'ten': 'An toàn và bảo vệ dữ liệu',
     'ngu_canh': 'an toàn trên mạng, mật khẩu, thông tin cá nhân',
     'tu_khoa': ['an toan', 'bao ve', 'mat khau', 'thong tin ca nhan', 'phan cung', 'ban phim',
                 'chuot', 'thiet bi', 'tu the ngoi', 'phong tranh', 'dau moi', 'dung luong'],
     'nls': 'An toàn: thực hành sử dụng thiết bị đúng cách; nhận biết thông tin cần bảo vệ; nêu một quy tắc an toàn phù hợp bài học.',
     'ai': 'Giáo viên chuẩn bị trước một phản hồi của AI về nội dung bài; học sinh đối chiếu với sách giáo khoa và chỉ ra chỗ AI có thể sai; không nhập thông tin cá nhân vào công cụ.'},
    {'ten': 'Khai thác dữ liệu và thông tin',
     'ngu_canh': 'thu thập số liệu, biểu đồ, thống kê, số liệu',
     'tu_khoa': ['tim kiem', 'thong tin', 'internet', 'du lieu', 'bang tinh', 'so lieu', 'thong ke',
                 'bieu do', 'tra cuu', 'tu lieu', 'ban do', 'khao sat', 'bao cao', 'bang bieu'],
     'nls': 'Khai thác dữ liệu và thông tin: đọc/tra cứu học liệu số do giáo viên lựa chọn; ghi lại số liệu cần dùng, đối chiếu hai nguồn và nêu nguồn tham khảo.',
     'ai': 'Học sinh dùng kết quả do AI gợi ý (giáo viên chuẩn bị trước) rồi đối chiếu với số liệu trong sách hoặc nguồn chính thống, nêu rõ điểm cần kiểm chứng.',
     'stem': 'Gợi ý STEM: thu thập số liệu ngoài lớp, ghi vào bảng tính rồi đọc kết quả bằng biểu đồ; nêu nhận xét và cách làm lại cho chính xác hơn.'},
    {'ten': 'Sáng tạo nội dung số',
     'ngu_canh': 'soạn thảo văn bản, trình chiếu, vẽ trên máy tính, tạo bài trình chiếu',
     'tu_khoa': ['ve tranh', 've hinh', 've so do', 'trinh chieu', 'van ban', 'thiep', 'thiet ke', 'soan thao', 'hinh anh',
                 'viet doan', 'ke chuyen', 'lam phim', 'poster', 'to roi', 'bao tuong', 'so tay'],
     'nls': 'Sáng tạo nội dung số: tạo một sản phẩm số ngắn minh họa nội dung bài; ghi nguồn hình ảnh/tư liệu và trình bày sản phẩm cho nhóm.',
     'stem': 'Gợi ý STEAM: thiết kế sản phẩm trực quan gắn nội dung bài; kết hợp công nghệ, bố cục/hình học và thẩm mỹ; trình bày lựa chọn thiết kế, nhận phản hồi và cải tiến.'},
    {'ten': 'Giải quyết vấn đề (thuật toán, lập trình)',
     'ngu_canh': 'lập trình, thuật toán, chương trình',
     'tu_khoa': ['lap trinh', 'scratch', 'robot', 'thuat toan', 'mo phong', 'mach dien', 'co cau',
                 'quy trinh', 'tung buoc', 'chuong trinh'],
     'nls': 'Giải quyết vấn đề: chia nhiệm vụ thành các bước, thử nghiệm chương trình hoặc mô hình, phát hiện lỗi và điều chỉnh dựa trên kết quả.',
     'stem': 'Gợi ý STEM: tổ chức nhiệm vụ thiết kế – chế tạo – thử nghiệm theo quy trình; ghi kết quả từng lần thử và cải tiến sản phẩm.'},
    {'ten': 'Giao tiếp và hợp tác trong môi trường số',
     'ngu_canh': 'thư điện tử, chia sẻ thông tin, trao đổi, hợp tác trực tuyến',
     'tu_khoa': ['hop tac', 'chia se', 'thu dien tu', 'giao tiep', 'thuyet trinh', 'thao luan',
                 'lam viec nhom', 'trao doi', 'dong vai', 'phong van', 'doc hieu', 'tap doc',
                 'ke lai', 'trinh bay', 'luyen noi', 'doc dien cam', 'nghe va noi', 'truyen',
                 'cau chuyen', 'bai tho', 'ca dao', 'listen', 'speak', 'unit',
                 'vocabulary', 'grammar', 'reading'],
     'nls': 'Giao tiếp và hợp tác trong môi trường số: chia sẻ sản phẩm trong nhóm dưới sự hướng dẫn của giáo viên, phản hồi lịch sự và không công khai dữ liệu cá nhân.',
     'ai': 'Học sinh soạn câu hỏi cho bạn cùng nhóm rồi dùng công cụ AI của lớp để gợi ý cách diễn đạt; cả nhóm kiểm tra lại xem gợi ý đó có phù hợp và trung thực không.'},
    {'ten': 'Ứng dụng trí tuệ nhân tạo',
     'ngu_canh': 'trí tuệ nhân tạo, chatbot, học máy',
     'tu_khoa': ['tri tue nhan tao', 'chatbot', 'tro ly ao', 'tu dong hoa', 'may moc thong minh',
                 'nhan dang', 'du doan'],
     'nls': 'Ứng dụng trí tuệ nhân tạo: nhận biết sản phẩm có AI, nêu việc AI làm được và việc AI có thể sai, kiểm tra lại trước khi dùng.',
     'ai': 'Giáo viên minh họa một phản hồi AI, học sinh đối chiếu với học liệu đã xác minh, chỉ ra điểm cần kiểm tra và nêu vì sao không đưa dữ liệu cá nhân vào câu lệnh.'},
    {'ten': 'Thí nghiệm, đo lường, quan sát',
     'ngu_canh': 'đo lường, đo dài, thí nghiệm, quan sát, ghi kết quả',
     'tu_khoa': ['thi nghiem', 'do dai', 'do nhiet', 'nhiet do', 'nang luong', 'tai che', 'trong cay',
                 'quan sat', 'nuoc', 'khong khi', 'dung dich', 'do luong', 'can nang', 'suc gio',
                 'thuc vat', 'dong vat', 'co the', 'suc khoe', 'thoi tiet', 'trai dat',
                 'cay xanh', 'cay coi', 'cham soc cay',
                 'mat troi', 'thuc pham', 'nam va vi khuan'],
     'nls': 'Giải quyết vấn đề: ghi kết quả đo/ quan sát vào bảng hoặc bảng tính số, so sánh các lần làm và giải thích chênh lệch bằng dữ liệu thu được.',
     'stem': 'Gợi ý STEM: tổ chức nhiệm vụ đo – thử – ghi kết quả – cải tiến gắn với bài học; dùng dụng cụ đơn giản và bảng ghi số liệu để học sinh giải thích kết quả.'},
    {'ten': 'Toán học với công cụ số',
     'ngu_canh': 'bài toán, phân số, phép cộng, diện tích',
     'tu_khoa': ['bang nhan', 'phep cong', 'phep tru', 'phep nhan', 'phep chia', 'phan so', 'ti so',
                 'phan tram', 'trung binh cong', 'hinh hoc', 'dien tich', 'the tich', 'so sanh so'],
     'nls': 'Giải quyết vấn đề: dùng công cụ số (bảng tính hoặc máy tính cầm tay) để kiểm tra lại kết quả, nêu các bước đã làm và chỗ dễ sai.',
     'ai': 'Học sinh tự làm bài trước, sau đó đối chiếu với kết quả do AI đưa ra (giáo viên chuẩn bị), tìm chỗ khác nhau và giải thích vì sao phải tự kiểm chứng.'},
    {'ten': 'Tra cứu tư liệu, bản đồ, mốc thời gian',
     'ngu_canh': 'bản đồ, tư liệu, nhân vật lịch sử, di tích',
     'tu_khoa': ['ban do', 'dien bien', 'chien dich', 'thoi ki', 'nhan vat lich su', 'di tich',
                 'vung mien', 'khi hau', 'dan so', 'van hoa', 'le hoi', 'dia hinh'],
     'nls': 'Khai thác dữ liệu và thông tin: tra cứu bản đồ/tư liệu số do giáo viên chọn, ghi lại dẫn chứng, so sánh hai nguồn và nêu nguồn đã dùng.',
     'ai': 'Học sinh kiểm chứng một thông tin do AI trả lời về bài học bằng sách giáo khoa hoặc tư liệu chính thống; ghi lại thông tin đúng và nêu lý do AI sai (nếu có).'},
    {'ten': 'Kĩ năng sống, quy tắc ứng xử',
     'ngu_canh': 'ứng xử, quy tắc, trung thực, kĩ năng sống',
     'tu_khoa': ['dao duc', 'ung xu', 'quy tac', 'trung thuc', 'tiet kiem', 'an toan giao thong',
                 'moi truong', 'gia dinh', 'ban be', 'cam xuc', 'ki nang song', 'hop tac xa hoi',
                 'loi hua', 'giu chu tin', 'doan ket', 'chia se voi ban'],
     'nls': 'An toàn: nêu việc nên và không nên làm trong tình huống của bài; phân biệt thông tin riêng tư với thông tin có thể chia sẻ.',
     'ai': 'Học sinh nghe/đọc một tình huống do AI tạo (giáo viên chuẩn bị), nhận xét việc làm nào đúng và nêu vì sao phải trung thực khi dùng kết quả của người khác.'},
    {'ten': 'Tạo hình, sản phẩm thủ công',
     'ngu_canh': 'vẽ tranh, tạo hình, trang trí, thủ công',
     'tu_khoa': ['ve tranh', 'trang tri', 'thu cong', 'gap hinh', 'cat dan', 'ban ve', 'tranh anh',
                 'mau sac', 'tiet kiem giay', 'do choi'],
     'nls': 'Sáng tạo nội dung số: chụp lại sản phẩm và tạo một trang giới thiệu ngắn (ảnh + vài dòng chữ) về cách làm sản phẩm của nhóm.',
     'stem': 'Gợi ý STEAM: thiết kế sản phẩm tạo hình theo yêu cầu bài học; thử với vật liệu khác nhau, nhận xét độ bền và tính thẩm mỹ rồi cải tiến.'},
]


def la_tin_hoc(subject):
    """Môn Tin học / Công nghệ thông tin: giữ cách tích hợp cho cả file như trước."""
    t = plain(subject or '')
    return 'tin hoc' in t or 'ict' in t or 'tin hoc nghe' in t


def _nhom_theo_tua(title):
    t = plain(title or '')
    hits = []
    for nhom in NHOM_TICH_HOP:
        khop = [k for k in nhom['tu_khoa'] if re.search(r'(?<![a-z0-9])' + re.escape(k) + r'(?![a-z0-9])', t)]
        if khop:
            hits.append((nhom, khop))
    return hits


def _khop_hien(title, k):
    """Cụm từ NGUYÊN VĂN trong tựa bài ứng với từ khoá đã khớp (giữ dấu tiếng Việt cho dễ đọc)."""
    tp = plain(title or '')
    m = re.search(r'(?<![a-z0-9])' + re.escape(k) + r'(?![a-z0-9])', tp)
    if not m:
        return k
    truoc = len(tp[:m.start()].split())
    tu = (title or '').split()
    return ' '.join(tu[truoc:truoc + len(m.group(0).split())]) or k


def cham_bai(title, subject='', grade=''):
    """Đọc tựa bài rồi suy luận: bài này có phù hợp để tích hợp hay không, và phù hợp nội dung nào."""
    hits = _nhom_theo_tua(title)
    khop = sorted({k for _, ks in hits for k in ks})
    diem = sum(3 * len(ks) for _, ks in hits) + (2 if hits else 0)
    if la_tin_hoc(subject):
        diem += 6                      # môn Tin học: bài nào cũng là bài học số
    return {'diem': diem, 'nhom': [n['ten'] for n, _ in hits], 'khop': khop,
            'hien': [_khop_hien(title, k) for k in khop]}


def ly_do_chon(title, cham):
    """Câu giải thích vì sao chọn bài này — in vào cột ghi chú để thầy/cô đối chiếu."""
    if cham.get('tin_hoc'):
        return ('Căn cứ tựa bài: môn Tin học nên mọi bài đều có thể tích hợp%s.' %
                ('' if not cham.get('khop') else
                 ' (tựa bài có “%s”)' % '”, “'.join((cham.get('hien') or cham['khop'])[:3])))
    if not cham.get('khop'):
        return 'Bài này chưa có từ khoá nào cho thấy cần tích hợp; để trống cho thầy/cô tự quyết định.'
    return ('Căn cứ tựa bài: có “%s” → nhóm nội dung “%s”.' %
            ('”, “'.join((cham.get('hien') or cham['khop'])[:4]),
             '”, “'.join(cham['nhom'][:2])))


def suggestions(title, grade='', subject=''):
    """Đọc tựa bài → suy luận nội dung tích hợp phù hợp (quy tắc rõ ràng, thầy/cô sửa lại được)."""
    result = {'digital': '', 'ai': '', 'stem': '', 'notes': ''}
    hits = _nhom_theo_tua(title)
    for nhom, _ in hits[:2]:
        if not result['digital']:
            result['digital'] = nhom['nls']
    for nhom, _ in hits[:2]:
        if not result['ai'] and nhom.get('ai'):
            result['ai'] = nhom['ai']
        if not result['stem'] and nhom.get('stem'):
            result['stem'] = nhom['stem']
    if any(result.values()):
        result['notes'] = (f'Đề xuất tham khảo cho lớp {grade}, môn {subject}. '
                           'GV xác nhận mục tiêu, học liệu và thời lượng; chưa tự gán mã chỉ báo.')
    return result


def extract_rows(doc):
    result = []
    for table in doc.tables:
        rows = table.rows
        header_index, columns = None, {}
        for ri, row in enumerate(rows[:5]):
            headers = [plain(c.text) for c in row.cells]
            for ci, value in enumerate(headers):
                if 'ten bai' in value or value.strip() in ('bai hoc', 'noi dung bai hoc'):
                    header_index = ri
                    columns['title'] = ci
            if header_index is not None:
                # Include the preceding merged header (e.g. Tuần) when needed.
                for ci, value in enumerate(headers):
                    prior = plain(rows[ri - 1].cells[ci].text) if ri > 0 and ci < len(rows[ri-1].cells) else ''
                    value = value + ' ' + prior
                    for key, terms in {'week': ['tuan'], 'topic': ['chu de', 'mach noi dung'],
                        'periods': ['thoi luong', 'so tiet', 'tiet hoc'], 'digital': ['nang luc so', 'nang luc so', 'nang luc so'],
                        'ai': ['tri tue nhan tao'], 'stem': ['stem', 'steam'], 'notes': ['ghi chu']}.items():
                        if any(x in value for x in terms):
                            columns.setdefault(key, ci)
                break
        if header_index is None:
            continue
        seen_cells = {}
        for row in rows[header_index+1:]:
            cells = row.cells
            item = {k: cells[i].text.strip() if i < len(cells) else '' for k,i in columns.items()}
            if not any(item.get(k) for k in ('title','week','periods','notes')) or 'ten bai' in plain(item.get('title','')):
                continue
            entry = {k: item.get(k,'') for k in ('week','topic','title','periods','digital','ai','stem','notes')}
            entry['_merge'] = {}
            for key, ci in columns.items():
                if ci < len(cells):
                    element = cells[ci]._tc
                    lookup = (key, element)
                    if lookup in seen_cells:
                        entry['_merge'][key] = seen_cells[lookup]
                    else:
                        seen_cells[lookup] = len(result)
            result.append(entry)
    if not result:
        for p in doc.paragraphs:
            if re.match(r'^(bài|bai|chủ đề)\s*\d+', p.text.strip(), re.I):
                result.append(dict(week='',topic='',title=p.text.strip(),periods='',digital='',ai='',stem='',notes=''))
    if not result:
        raise ValueError('Chưa nhận diện được bảng phân phối chương trình. File cần cột “Tên bài học” hoặc các dòng bắt đầu bằng “Bài 1…”.')
    if len(result) > 180:
        raise ValueError('Mỗi lần hỗ trợ tối đa 180 dòng bài học. Hãy chia file theo học kỳ.')
    return result


def preview(data, mode, grade, subject, chon=None, so_bai=None):
    """chon: tập nội dung cần đề xuất — 'digital' (năng lực số), 'ai' (giáo dục AI), 'stem'.

    so_bai (M7): số bài tối đa cần tích hợp ở phân phối chương trình. Để trống thì với môn ngoài
    Tin học hệ thống chỉ chọn những bài có tựa bài thật sự phù hợp; môn Tin học giữ nguyên cách
    tích hợp cho cả file. Trong cả hai trường hợp, hệ thống đọc tựa bài rồi mới suy luận nội dung.
    """
    chon = set(chon or ('digital', 'ai', 'stem'))
    tin_hoc = la_tin_hoc(subject)
    doc = read_word(data)
    if mode == 'ppct':
        rows = extract_rows(doc)
        for row in rows:
            g = cham_bai(row['title'], subject, grade)
            row['diem_phu_hop'] = g['diem']
            row['nhom_noi_dung'] = g['nhom']
            row['_khop'] = g['khop']
            row['_khop_hien'] = g['hien']
        if so_bai is None:
            # để trống: môn Tin học tích hợp cả file, môn khác chỉ những bài tựa bài phù hợp
            chon_dong = (set(range(len(rows))) if tin_hoc else
                         {i for i, r in enumerate(rows) if r['diem_phu_hop'] > 0})
        else:
            xep = sorted(range(len(rows)), key=lambda i: (-rows[i]['diem_phu_hop'], i))
            xep = [i for i in xep if rows[i]['diem_phu_hop'] > 0][:max(0, int(so_bai))]
            chon_dong = set(xep)
    else:
        text = '\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
        if not text.strip():
            raise ValueError('Không tìm thấy văn bản trong file; file ảnh quét chưa được hỗ trợ.')
        title = next((p.text.strip() for p in doc.paragraphs if re.search(r'(bài|chủ đề|tên bài)', p.text, re.I)), '')
        title = title or next((p.text.strip() for p in doc.paragraphs if p.text.strip()), 'Giáo án')
        rows = [dict(week='', topic='', title=title[:500], periods='', digital='', ai='', stem='', notes='')]
        for row in rows:
            g = cham_bai(row['title'], subject, grade)
            row['diem_phu_hop'] = g['diem']
            row['nhom_noi_dung'] = g['nhom']
            row['_khop'] = g['khop']
            row['_khop_hien'] = g['hien']
        chon_dong = {0}
    from .framework_match import match
    if 'digital' not in chon:
        text = ''                      # không đề xuất năng lực số: bỏ luôn phần đối chiếu chỉ báo
    for i, row in enumerate(rows):
        row['original'] = {k: row[k] for k in ('digital','ai','stem','notes')}
        row['de_xuat_tich_hop'] = i in chon_dong
        row['ly_do'] = ''
        row['references'] = []
        row['ma_ai'] = []
        row['canh_bao_ai'] = []
        row['chon'] = sorted(chon)
        if i not in chon_dong:
            # (M7) bài không nằm trong số bài cần tích hợp: để trống, thầy/cô tự nhập nếu muốn
            row['canh_bao_ai'] = ['Bài “%s” chưa được đề xuất tích hợp; thầy/cô có thể tự nhập nếu '
                                  'thấy phù hợp với lớp mình.' % row['title'][:80]]
            continue
        proposal = suggestions(row['title'], grade, subject)
        _goc_notes = row['original']['notes'].strip()
        _goi_y = (proposal.get('digital') or '').strip()      # phần diễn giải -> chỉ nằm ở cột Ghi chú
        _goi_y_ai = (proposal.get('ai') or '').strip()
        # (M8) cột tích hợp chỉ ghi MÃ TIÊU CHÍ; STEM chưa có hệ mã nên vẫn gợi ý bằng nội dung
        for key in ('stem',):
            if key in chon and not row[key]:
                row[key] = proposal.get(key, '')
        if 'digital' not in chon:
            row['digital'] = row['original']['digital']
            row['digital_cu'] = []
        if 'stem' not in chon:
            row['stem'] = row['original']['stem']
        _notes = []
        # ---- mã giáo dục AI theo Quyết định 2422/QĐ-BGDĐT (không tự đặt mã) ----
        if 'ai' in chon and not row['original']['ai']:
            ds_ma, canh_bao_ai = AIGD.goi_y_ma(grade, ten_bai=row['title'],
                                               van_ban=row.get('topic', ''), mon=subject)
            row['ai'] = '\n'.join(x['ma'] for x in ds_ma)
            row['ma_ai'] = [{'ma': x['ma'], 'yccd': x['yccd'], 'noi_dung': x.get('noi_dung', ''),
                             'chu_de_ten': x['chu_de_ten'], 'mach_ten': x['mach_ten'],
                             'mo_rong': x['mo_rong'], 'khop': x.get('khop', []),
                             'nguon': AIGD.NGUON_MA} for x in ds_ma]
            row['canh_bao_ai'] = canh_bao_ai
            if ds_ma:
                _notes.append('Giáo dục AI (Quyết định 2422/QĐ-BGDĐT + Công văn 5588/BGDĐT-GDPT): ' +
                              '; '.join(x['yccd'][:110] for x in ds_ma) +
                              '\nGiáo viên rà soát mã và nội dung AI trước khi dùng.')
            elif _goi_y_ai:
                _notes.append('Hoạt động gợi ý (AI): ' + _goi_y_ai)
        elif 'ai' not in chon:
            row['ai'] = row['original']['ai']
        # ---- (M8) đọc tựa bài → lấy MÃ TIÊU CHÍ năng lực số phù hợp nhất trong khung ----
        references = match(text if mode == 'lesson' else row['title'], grade)
        _da_co = {x['code'] for x in references}
        for _nhom, _ks in _nhom_theo_tua(row['title'])[:2]:
            if not _nhom.get('ngu_canh'):
                continue
            for x in match(_nhom['ngu_canh'], grade, toi_da=2):
                if x['code'] in _da_co:
                    continue
                _da_co.add(x['code'])
                _x = dict(x)
                _x['evidence'] = 'tựa bài: ' + '”, “'.join(_khop_hien(row['title'], k) for k in _ks[:2])
                references.append(_x)
        references = references[:3]
        row['references'] = references
        if 'digital' in chon and not row['original']['digital']:
            if references:
                row['digital'] = '\n'.join(x['code'] for x in references)
                _notes += [x for x in dict.fromkeys(y['activity'] for y in references if y.get('activity'))]
            elif _goi_y:
                _notes.append('Chưa có mã chỉ báo năng lực số nào thật sát bài này trong khung — '
                              'thầy/cô tự chọn mã của lớp; hệ thống không tự đặt mã.')
            if _goi_y:
                _notes.append('Nội dung gợi ý: ' + _goi_y)
        _notes = ([_goc_notes] if _goc_notes else []) + _notes
        if _notes:
            _notes.append('Đề xuất tự động theo khung; giáo viên rà soát trước khi sử dụng.')
        row['notes'] = '\n'.join(_notes)
        # ---- (M7) ghi rõ căn cứ chọn bài để thầy/cô đối chiếu ----
        row['ly_do'] = ly_do_chon(row['title'], {'khop': row['_khop'], 'hien': row['_khop_hien'],
                                                'nhom': row['nhom_noi_dung'], 'tin_hoc': tin_hoc})
        row['notes'] = (row['ly_do'] + '\n' + row['notes']) if row['notes'] else row['ly_do']
    return rows


def export(data, mode, rows, grade, subject):
    if mode == 'ppct':
        doc = Document(TEMPLATE)
        doc.paragraphs[0].text = f'Môn học, hoạt động giáo dục: {subject} — Lớp {grade}'
        doc.paragraphs[1].text = 'Thời lượng và tuần học: theo các dòng phân phối chương trình bên dưới; giáo viên rà soát trước khi sử dụng.'
        table = doc.tables[0]
        while len(table.rows) > 2:
            table._tbl.remove(table.rows[-1]._tr)
        for item in rows:
            cells = table.add_row().cells
            for cell, key in zip(cells, ('week','topic','title','periods','digital','ai','stem','notes')):
                cell.text = item.get(key, '')
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(11)
                _to_mau(cell, key, item)      # đỏ: năng lực số, xanh dương: giáo dục AI
        keys = ('week','topic','title','periods','digital','ai','stem','notes')
        # Recreate source vertical merges only if the teacher kept matching values.
        groups = {}
        for ri, item in enumerate(rows):
            for key, first in item.get('_merge', {}).items():
                if key in keys and 0 <= first < ri:
                    groups.setdefault((key, first), []).append(ri)
        for (key, first), rest in groups.items():
            last = max(rest)
            if rest == list(range(first+1,last+1)) and all(rows[x].get(key,'') == rows[first].get(key,'') for x in rest):
                ci = keys.index(key)
                cell = table.cell(first+2, ci).merge(table.cell(last+2, ci))
                cell.text = rows[first].get(key,'')
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(11)
                _to_mau(cell, key, rows[first])
    else:
        doc = read_word(data)
        # Append a reviewed integration section so original lesson tables/runs remain intact.
        doc.add_heading('TÍCH HỢP NĂNG LỰC SỐ, AI VÀ STEM/STEAM', level=1)
        doc.add_paragraph(f'Môn {subject} — Lớp {grade}')
        for item in rows:
            for label, key in (('Năng lực số','digital'),('Trí tuệ nhân tạo (AI)','ai'),('STEM/STEAM','stem'),('Tổ chức và đánh giá','notes')):
                if item.get(key):
                    doc.add_heading(label, level=2)
                    p = doc.add_paragraph()
                    run = p.add_run(item[key])
                    mau = MAU_COT.get(key)
                    if mau is not None and item.get('_themmoi', {}).get(key):
                        run.font.color.rgb = mau
    doc.add_paragraph('Tham chiếu Thông tư 02/2025/TT-BGDĐT và hướng dẫn theo khối lớp tại Công văn 3456/BGDĐT-GDPT (mã năng lực số). Mã giáo dục AI lấy theo Quyết định 2422/QĐ-BGDĐT ngày 18/8/2026 (Khung nội dung giáo dục AI cho học sinh phổ thông; quy ước mã [Lớp].[Mã chủ đề].[Số thứ tự], nội dung mở rộng thêm tiền tố “MR”) và hướng dẫn triển khai tại Công văn 5588/BGDĐT-GDPT ngày 19/8/2026. Các mã và hoạt động tích hợp cần được giáo viên rà soát theo thực tế lớp học.')
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
