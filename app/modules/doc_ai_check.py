"""Dò dấu hiệu văn bản do MÁY (AI) viết — chỉ mang tính THAM KHẢO, không phải bằng chứng.

Cách làm: trích các đặc trưng văn phong (nhịp câu, độ đều, mật độ từ nối, cụm lặp, dấu câu,
từ ngữ khẩu ngữ, số liệu cụ thể...) rồi cho một mô hình hồi quy logistic đã huấn luyện
chấm điểm từng đặc trưng. Trọng số nằm trong `app/assets/ml/ml_do_van_ai.json`,
đọc bằng numpy nên máy chủ yếu (1 nhân / 1 GB) vẫn chạy được.

NÓI THẲNG VỀ GIỚI HẠN (bắt buộc đọc trước khi dùng kết quả):
  · Đây KHÔNG phải mô hình ngôn ngữ lớn, không đọc hiểu nội dung. Nó chỉ đếm và đo văn phong.
  · Mọi công cụ dò văn bản AI trên thị trường đều sai nhiều, kể cả của các công ty lớn.
  · Văn bản ngắn (< 120 từ) thì gần như không kết luận được gì.
  · Văn bản hành chính, báo cáo, sách giáo khoa do NGƯỜI viết rất dễ bị chấm nhầm là AI
    vì văn phong vốn đều đặn, nhiều từ nối.
  · Vì vậy: KHÔNG dùng kết quả này để buộc tội, trừ điểm hay kết luận về một học sinh.
    Chỉ dùng như một gợi ý để giáo viên hỏi lại, xem lại, hoặc đối chiếu với quá trình học.
"""
import io
import json
import os
import re

import numpy as np

# ------------------------------------------------------------------ từ điển dấu hiệu
TU_NOI = (
    "do đó", "vì vậy", "bởi vậy", "hơn nữa", "ngoài ra", "bên cạnh đó", "mặt khác",
    "tóm lại", "nhìn chung", "tổng kết lại", "có thể thấy rằng", "điều này cho thấy",
    "đáng chú ý là", "đặc biệt là", "chính vì thế", "song song đó", "theo đó",
    "trong bối cảnh", "trên thực tế", "nói cách khác", "đối với việc", "thông qua việc",
    "một cách toàn diện", "đóng vai trò quan trọng", "không chỉ", "mà còn",
)
DAU_MO_DAN = ("mở rộng", "tối ưu", "toàn diện", "hiệu quả", "nâng cao", "phát triển",
                 "trải nghiệm", "giải pháp", "xu hướng", "tiềm năng", "đột phá", "linh hoạt")
KHAU_NGU = ("em", "mình", "tớ", "con", "ạ", "nhé", "nha", "hả", "hơi", "quá", "lắm",
            "thì", "mà", "chứ", "vậy", "ơi", "nè", "ừ", "ờm", "ha")
MARKDOWN = ("**", "__", "##", "```", "- ", "* ", "1. ", "2. ", "3. ", "• ")
EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF]")
TU = re.compile(r"[0-9A-Za-zÀ-ỹĐđ]+", re.UNICODE)
CAU = re.compile(r"[^.!?…\n]+[.!?…]*")

DAC_TRUNG = [
    "cau_lech", "cau_deu", "cau_dai", "cau_ngan",          # nhịp điệu câu
    "ttr", "tu_dai", "lap_cum", "lap_tu_gan",              # từ vựng
    "tu_noi", "mo_dau_tu_noi", "dau_mo_dan",               # từ nối kiểu máy
    "khau_ngu", "dai_tu", "so_lieu",                        # dấu hiệu người viết
    "dau_cau", "phay", "bullet", "markdown", "emoji",       # hình thức
    "doan_deu", "cau_truc_lap", "muot",                     # độ đều / độ "mượt"
]

_ML = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "ml")
_CACHE = {}


# ------------------------------------------------------------------ trích văn bản
def tach_doan(text):
    """Tách thành các đoạn không rỗng."""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    ra = []
    for p in re.split(r"\n\s*\n|\n(?=\s*(?:[IVX]+\.|[0-9]+\.|[•\-*]))", text):
        p = re.sub(r"[ \t]+", " ", p).strip()
        if p:
            ra.append(p)
    return ra


def doc_docx(blob):
    """Đọc file Word: lấy cả chữ trong bảng (bảng là chỗ dễ sót nhất)."""
    from docx import Document
    d = Document(io.BytesIO(blob))
    phan = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    if p.text.strip():
                        phan.append(p.text)
    return phan


def doc_pdf(blob):
    """Đọc file PDF (chỉ lấy được chữ; PDF là ảnh quét thì báo rõ là không đọc được)."""
    try:
        from pypdf import PdfReader
    except Exception:
        return [], "Máy chủ chưa có thư viện đọc PDF."
    try:
        r = PdfReader(io.BytesIO(blob))
        phan = []
        for trang in r.pages:
            t = (trang.extract_text() or "").strip()
            if t:
                phan.extend(tach_doan(t))
        if not phan:
            return [], ("PDF này không có lớp chữ (có thể là ảnh quét). "
                        "Thầy/cô hãy dùng bản Word hoặc dán văn bản vào ô bên cạnh.")
        return phan, ""
    except Exception as e:
        return [], "Không đọc được PDF này (%s)." % type(e).__name__


# ------------------------------------------------------------------ đặc trưng
def _cau_tu(doan):
    cau, tu = [], []
    for p in doan:
        for m in CAU.finditer(p):
            c = m.group(0).strip()
            if c:
                cau.append(c)
                tu.extend(TU.findall(c.lower()))
    return cau, tu


def dac_trung(doan):
    """Tính vector đặc trưng. Trả về dict tên -> giá trị (số thực)."""
    cau, tu = _cau_tu(doan)
    text = " ".join(doan)
    thap = text.lower()
    dai_cau = [len(TU.findall(c.lower())) for c in cau] or [0]
    tb = float(np.mean(dai_cau)) if dai_cau else 0.0
    lech = float(np.std(dai_cau)) if len(dai_cau) > 1 else 0.0
    n_tu = max(len(tu), 1)

    def _tl(ds):
        return sum(1 for x in ds if x) / float(len(ds)) if ds else 0.0

    ttr = len(set(tu[:400])) / float(max(len(tu[:400]), 1))
    cum = [" ".join(tu[i:i + 4]) for i in range(max(len(tu) - 3, 0))]
    lap = 1.0 - (len(set(cum)) / float(len(cum))) if cum else 0.0
    lap_gan = sum(1 for i in range(1, len(tu)) if tu[i] == tu[i - 1]) / float(n_tu)
    dai_doan = [len(TU.findall(p.lower())) for p in doan] or [0]
    tb_doan = float(np.mean(dai_doan)) if dai_doan else 0.0
    doan_deu = 0.0
    if tb_doan > 0 and len(dai_doan) > 1:
        doan_deu = sum(1 for x in dai_doan if abs(x - tb_doan) <= 0.2 * tb_doan) / float(len(dai_doan))
    # câu mở đầu giống nhau
    mo = [" ".join(TU.findall(c.lower())[:2]) for c in cau if TU.findall(c.lower())]
    cau_truc_lap = 1.0 - (len(set(mo)) / float(len(mo))) if len(mo) > 1 else 0.0
    # các câu mở đầu bằng từ nối
    dau_noi = sum(1 for c in cau if any(c.lower().lstrip("“\"(- ").startswith(t) for t in TU_NOI))
    so_phay = text.count(",") + text.count(";")

    f = {
        "cau_lech": lech / tb if tb else 0.0,
        "cau_deu": _tl([1 if tb and abs(x - tb) <= 0.15 * tb else 0 for x in dai_cau]),
        "cau_dai": _tl([1 if x > 28 else 0 for x in dai_cau]),
        "cau_ngan": _tl([1 if x < 6 else 0 for x in dai_cau]),
        "ttr": ttr,
        "tu_dai": float(np.mean([len(w) for w in tu])) if tu else 0.0,
        "lap_cum": lap,
        "lap_tu_gan": lap_gan,
        "tu_noi": sum(thap.count(t) for t in TU_NOI) * 100.0 / n_tu,
        "mo_dau_tu_noi": dau_noi / float(len(cau)) if cau else 0.0,
        "dau_mo_dan": sum(thap.count(t) for t in DAU_MO_DAN) * 100.0 / n_tu,
        "khau_ngu": sum(1 for w in tu if w in KHAU_NGU) * 100.0 / n_tu,
        "dai_tu": sum(1 for w in tu if w in ("em", "mình", "tớ", "con", "bạn")) * 100.0 / n_tu,
        "so_lieu": sum(1 for ch in text if ch.isdigit()) * 100.0 / n_tu,
        "dau_cau": (text.count(".") + text.count("!") + text.count("?") + so_phay) * 100.0 / n_tu,
        "phay": so_phay * 100.0 / n_tu,
        "bullet": sum(thap.count(m) for m in ("•", "- ", "* ")) * 100.0 / n_tu,
        "markdown": float(sum(thap.count(m) for m in MARKDOWN)),
        "emoji": float(len(EMOJI.findall(text))),
        "doan_deu": doan_deu,
        "cau_truc_lap": cau_truc_lap,
        "muot": _do_muot(doan),
    }
    return {k: float(f.get(k, 0.0)) for k in DAC_TRUNG}


def _lm():
    """Mô hình ngôn ngữ ký tự (trigram) huấn luyện trên văn bản người viết."""
    if "lm" not in _CACHE:
        try:
            with open(os.path.join(_ML, "lm_chu_viet.json"), encoding="utf-8") as f:
                d = json.load(f)
            _CACHE["lm"] = (d.get("dem", {}), d.get("tong", 0), d.get("so_ky_tu", 0))
        except Exception:
            _CACHE["lm"] = ({}, 0, 0)
    return _CACHE["lm"]


def _do_muot(doan):
    """Độ 'mượt' = trung bình log xác suất trigram ký tự. Văn bản càng giống văn mẫu càng cao."""
    dem, tong, n_ky = _lm()
    if not dem or not n_ky:
        return 0.0
    s = " ".join(doan).lower()
    s = re.sub(r"\s+", " ", s)
    if len(s) < 60:
        return 0.0
    diem = []
    for i in range(2, len(s)):
        tri = s[i - 2:i + 1]
        hai = s[i - 2:i + 1][:2]
        c = dem.get(tri, 0)
        t = dem.get(hai, 0)
        diem.append(np.log((c + 0.4) / (t + 0.4 * 95.0)))
    return float(np.mean(diem)) if diem else 0.0


# ------------------------------------------------------------------ mô hình
def _mh():
    if "mh" not in _CACHE:
        try:
            with open(os.path.join(_ML, "ml_do_van_ai.json"), encoding="utf-8") as f:
                _CACHE["mh"] = json.load(f)
        except Exception:
            _CACHE["mh"] = None
    return _CACHE["mh"]


def co_mo_hinh():
    return bool(_mh() and _mh().get("trong_so"))


def thong_tin():
    m = _mh() or {}
    return {"co_mo_hinh": bool(m), "phien_ban": m.get("phien_ban", ""),
            "huan_luyen_luc": m.get("huan_luyen_luc", ""),
            "chi_so": m.get("chi_so", {}), "chi_so_kho": m.get("chi_so_kho", {}),
            "chi_so_thuc_te": m.get("chi_so_thuc_te", {}), "chi_so_giao_an": m.get("chi_so_giao_an", {}),
            "du_lieu": m.get("du_lieu", {}),
            "so_dac_trung": len(m.get("trong_so") or {}), "canh_bao": m.get("canh_bao", "")}


def _logit(ft):
    m = _mh()
    if not m:
        return None, 0.0
    w = m["trong_so"]
    tb = m.get("trung_binh", {})
    lc = m.get("do_lech", {})
    z = float(m.get("he_so_tu_do", 0.0))
    for k in DAC_TRUNG:
        x = float(ft.get(k, 0.0))
        sd = lc.get(k, 1.0) or 1.0
        z += float(w.get(k, 0.0)) * ((x - tb.get(k, 0.0)) / sd)
    return z, 1.0 / (1.0 + np.exp(-z))


MUC = ((35, "rất ít dấu hiệu máy viết", "#059669"),
       (65, "chưa kết luận được", "#d97706"),
       (101, "có khá nhiều dấu hiệu máy viết", "#e11d48"))


def _muc(diem):
    for nguong, ten, mau in MUC:
        if diem < nguong:
            return ten, mau
    return MUC[-1][1], MUC[-1][2]


LY_DO = {
    "tu_noi": "mật độ từ nối kiểu văn mẫu (do đó, hơn nữa, nhìn chung…)",
    "mo_dau_tu_noi": "tỉ lệ câu mở đầu bằng từ nối",
    "dau_mo_dan": "nhiều từ ngữ kiểu văn nghị luận mẫu (tối ưu, toàn diện, nâng cao…)",
    "cau_deu": "độ dài các câu rất đều nhau",
    "cau_lech": "nhịp câu ít thay đổi",
    "doan_deu": "các đoạn dài ngắn gần như bằng nhau",
    "cau_truc_lap": "nhiều câu mở đầu giống nhau",
    "lap_cum": "lặp lại cùng những cụm từ",
    "cau_dai": "nhiều câu dài",
    "muot": "câu chữ trôi rất mượt, giống văn mẫu",
    "khau_ngu": "ít từ khẩu ngữ đời thường",
    "dai_tu": "ít cách gọi gần gũi (em, mình, con…)",
    "so_lieu": "ít chi tiết số liệu cụ thể",
    "lap_tu_gan": "có chỗ lặp từ liền nhau (nét của người viết vội)",
    "bullet": "nhiều dòng gạch đầu dòng",
    "markdown": "có ký hiệu định dạng kiểu ** __ ## của công cụ soạn thảo",
    "emoji": "có biểu tượng cảm xúc trong văn bản",
    "ttr": "vốn từ lặp lại nhiều",
    "cau_ngan": "câu quá ngắn, rời rạc",
    "dau_cau": "mật độ dấu câu cao",
    "phay": "nhiều dấu phẩy, câu dài nhiều vế",
    "tu_dai": "từ ngữ dài, nhiều từ Hán Việt",
}


def cham(doan, chi_tiet=True):
    """Chấm điểm. Trả về điểm 0–100 (càng cao càng nhiều dấu hiệu máy viết) + lý do."""
    doan = [p for p in (doan or []) if p and p.strip()]
    so_tu = len(TU.findall(" ".join(doan).lower()))
    ft = dac_trung(doan)
    z, p = _logit(ft)
    if z is None:
        return {"co_mo_hinh": False, "diem": None, "muc": "chưa có mô hình",
                "mau": "#64748b", "so_tu": so_tu, "ly_do": [], "phan_doi": [],
                "tung_doan": [], "canh_bao": ""}
    diem = round(100.0 * float(p), 1)
    ten, mau = _muc(diem)

    # lý do: đặc trưng nào đẩy điểm lên mạnh nhất
    m = _mh()
    w, tb, lc = m["trong_so"], m.get("trung_binh", {}), m.get("do_lech", {})
    gop = []
    for k in DAC_TRUNG:
        x = ft.get(k, 0.0)
        sd = lc.get(k, 1.0) or 1.0
        g = float(w.get(k, 0.0)) * ((x - tb.get(k, 0.0)) / sd)
        gop.append((g, k, x, tb.get(k, 0.0)))
    gop.sort(reverse=True)
    ly_do = []
    for g, k, x, tbv in gop[:6]:
        if g <= 0.25:
            break
        if k in ("bullet", "markdown", "emoji"):
            ly_do.append("%s (%d lần)" % (LY_DO[k], int(x)))
        elif k in ("tu_noi", "dau_mo_dan"):
            ly_do.append("%s: %.1f lần/100 từ" % (LY_DO[k], x))
        elif k in ("cau_deu", "cau_dai", "cau_ngan", "mo_dau_tu_noi", "cau_truc_lap", "doan_deu"):
            ly_do.append("%s (%.0f%%)" % (LY_DO[k], x * 100))
        else:
            ly_do.append("%s (%.2f, mức thường %.2f)" % (LY_DO[k], x, tbv))
    phan_doi = []
    for g, k, x, tbv in reversed(gop):
        if g >= -0.25 or k not in LY_DO:
            continue
        phan_doi.append(LY_DO[k])
        if len(phan_doi) >= 4:
            break
    out = {"co_mo_hinh": True, "diem": diem, "muc": ten, "mau": mau, "so_tu": so_tu,
           "ly_do": ly_do, "phan_doi": phan_doi,
           "dac_trung": {k: round(ft[k], 3) for k in DAC_TRUNG},
           "canh_bao": ("Văn bản quá ngắn (dưới 120 từ) — kết quả gần như không có ý nghĩa."
                        if so_tu < 120 else "")}
    if chi_tiet:
        out["tung_doan"] = _tung_doan(doan)
    return out


def _tung_doan(doan):
    """Chấm từng đoạn để giáo viên thấy chỗ nào đáng xem lại (đoạn < 40 từ thì bỏ qua)."""
    ra = []
    for i, p in enumerate(doan):
        if len(TU.findall(p.lower())) < 40:
            ra.append({"i": i, "diem": None})
            continue
        z, pr = _logit(dac_trung([p]))
        ra.append({"i": i, "diem": round(100.0 * float(pr), 1) if pr is not None else None})
    return ra
