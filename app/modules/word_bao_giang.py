"""Xuất lịch báo giảng 1 tuần ra file Word (.docx) — khổ A4 ngang, cùng mẫu với bản PDF.

Ưu điểm so với PDF: giáo viên mở bằng Word/WPS chỉnh sửa thêm trước khi in hoặc nộp tổ.
"""
import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

THU_NAME = {2: "Thứ Hai", 3: "Thứ Ba", 4: "Thứ Tư", 5: "Thứ Năm",
            6: "Thứ Sáu", 7: "Thứ Bảy", 8: "Chủ Nhật"}

FONT = "Times New Roman"          # font chuẩn văn bản hành chính Việt Nam
HDR_BG = "D6EFE6"                 # xanh lá nhạt, đồng bộ tông website
NGHI_BG = "FFE9EC"                # hồng nhạt cho dòng nghỉ lễ


def _set_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    # bắt buộc để Word áp font cho cả ký tự tiếng Việt
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)


def _para(container, text="", size=11, bold=False, align=None, italic=False,
          space_after=0, color=None):
    p = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    _set_font(r, size, bold, italic, color)
    return p


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000")
        borders.append(e)
    tblPr.append(borders)


def _cell(cell, text, size=10, bold=False, align=None, color=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        _set_font(r, size, bold, italic, color)
    return cell


def build_docx(meta, rows):
    """meta: dict(truong, to, giao_vien, mon, tuan, tu_ngay, den_ngay, noi_dung_khac, nghi, ...)
    rows: list dict(thu, buoi, tiet, lop, mon, tiet_pp, ten_bai, ghi_chu, ngay, nghi)"""
    doc = Document()

    # ---- khổ A4 ngang, lề hẹp ----
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.0)

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)

    C = WD_ALIGN_PARAGRAPH.CENTER
    L = WD_ALIGN_PARAGRAPH.LEFT

    # ---- phần đầu: trường / tổ | tiêu đề ----
    head = doc.add_table(rows=2, cols=2)
    head.autofit = False
    head.columns[0].width = Cm(7.5)
    head.columns[1].width = Cm(19.8)
    _cell(head.cell(0, 0), (meta.get("truong") or "").upper(), 10, True, C)
    _cell(head.cell(0, 1), "LỊCH BÁO GIẢNG", 16, True, C)
    _cell(head.cell(1, 0), "Tổ: " + (meta.get("to") or ""), 10, False, C)
    _cell(head.cell(1, 1),
          f"TUẦN {meta.get('tuan','')} (Từ ngày {meta.get('tu_ngay','')} "
          f"đến ngày {meta.get('den_ngay','')})"
          + ((" · TKB áp dụng từ tuần %s" % meta["ap_dung_tu_tuan"]) if meta.get("ap_dung_tu_tuan") else ""),
          11, False, C)

    _para(doc, "", 4, space_after=2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    _set_font(p.add_run("Họ và tên giáo viên: "), 11)
    _set_font(p.add_run(meta.get("giao_vien", "")), 11, bold=True)
    if meta.get("mon"):
        _set_font(p.add_run("          Môn dạy: "), 11)
        _set_font(p.add_run(meta.get("mon", "")), 11, bold=True)

    # ---- bảng chính ----
    hdr = ["Thứ/Ngày", "Buổi", "Tiết", "Lớp", "Môn", "Tiết PPCT", "Tên bài dạy", "Đồ dùng / Ghi chú"]
    widths = [Cm(2.7), Cm(1.7), Cm(1.2), Cm(1.8), Cm(3.0), Cm(1.8), Cm(9.0), Cm(6.1)]

    from .pdf_bao_giang import gop_hang
    ordered = gop_hang(rows)
    table = doc.add_table(rows=1 + max(len(ordered), 1), cols=len(hdr))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _borders(table)

    for i, h in enumerate(hdr):
        c = table.cell(0, i)
        _cell(c, h, 10, True, C)
        _shade(c, HDR_BG)

    if not ordered:
        merged = table.cell(1, 0).merge(table.cell(1, len(hdr) - 1))
        _cell(merged, "(Chưa có dữ liệu thời khoá biểu)", 10, False, C, italic=True)
    else:
        for ri, r in enumerate(ordered, start=1):
            thu = int(r.get("thu") or 9)
            label = THU_NAME.get(thu, "")
            if r.get("ngay"):
                label += "\n" + r["ngay"]
            is_nghi = bool(r.get("nghi"))
            ph = (r.get("phong") or "").strip()
            gc = (r.get("ghi_chu") or "").strip()
            if ph and ph not in gc:
                gc = (ph + " · " + gc) if gc else ph
            vals = [label, r.get("buoi", ""), r.get("tiet", ""), r.get("lop", ""), r.get("mon", ""),
                    r.get("tiet_pp", ""), r.get("ten_bai", ""), gc]
            aligns = [C, C, C, C, L, C, L, L]
            for ci, (v, al) in enumerate(zip(vals, aligns)):
                bold = (ci == 0) or (is_nghi and ci == 6)
                color = "BE123C" if (is_nghi and ci == 6) else None
                _cell(table.cell(ri, ci), "" if v is None else str(v), 10, bold, al, color)
            if is_nghi:
                for ci in range(1, len(hdr)):
                    _shade(table.cell(ri, ci), NGHI_BG)

        # gộp ô cột Thứ/Ngày (cùng ngày) và cột Buổi (cùng Sáng / Chiều)
        for idx, r in enumerate(ordered):
            ri = idx + 1
            rs = int(r.get("rowspan_thu") or 0)
            if rs > 1:
                table.cell(ri, 0).merge(table.cell(ri + rs - 1, 0))
            rb = int(r.get("rowspan_buoi") or 0)
            if rb > 1:
                table.cell(ri, 1).merge(table.cell(ri + rb - 1, 1))

    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    # ---- ghi chú ----
    if meta.get("nghi"):
        _para(doc, "", 4, space_after=2)
        p = doc.add_paragraph()
        _set_font(p.add_run("Ghi chú: "), 10, bold=True)
        _set_font(p.add_run("Tuần này trùng kỳ nghỉ - " + meta["nghi"]), 10)
    if meta.get("noi_dung_khac"):
        p = doc.add_paragraph()
        _set_font(p.add_run("Nội dung công tác khác: "), 10, bold=True)
        _set_font(p.add_run(meta["noi_dung_khac"]), 10)

    # ---- phần ký ----
    _para(doc, "", 6, space_after=4)
    sign = doc.add_table(rows=2, cols=2)
    sign.autofit = False
    sign.columns[0].width = Cm(13.6)
    sign.columns[1].width = Cm(13.7)
    _cell(sign.cell(0, 0), "DUYỆT CỦA TỔ TRƯỞNG CM", 10, True, C)
    _cell(sign.cell(0, 1),
          f"{meta.get('dia_danh','')}, ngày {meta.get('ngay_ky','....')} "
          f"tháng {meta.get('thang_ky','....')} năm {meta.get('nam_ky','......')}", 10, False, C)
    _cell(sign.cell(1, 0), "", 10)
    c = sign.cell(1, 1)
    _cell(c, "GIÁO VIÊN", 10, True, C)
    _para(c, "", 10, space_after=0)
    _para(c, "", 10, space_after=0)
    _para(c, meta.get("giao_vien", ""), 10, bold=True, align=C)

    # (LOGO-B) dòng thương hiệu ở đầu/chân trang
    try:
        from . import thuong_hieu as T_HIEU
        T_HIEU.gan(doc, tieu_de="Lịch báo giảng — %s" % (meta.get("mon", "") or ""))
    except Exception:
        pass
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
