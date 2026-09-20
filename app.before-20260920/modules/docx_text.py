"""Đọc và sửa văn bản trong file .docx mà GIỮ NGUYÊN định dạng gốc.

Kỹ thuật: mỗi đoạn (paragraph) gồm nhiều "run" có định dạng riêng.
Khi thay chữ, ta ánh xạ vị trí ký tự về đúng run chứa nó rồi chỉ sửa text của run đó,
nhờ vậy font/cỡ chữ/in đậm/màu... không bị mất.
Hỗ trợ cả văn bản trong bảng (table) và header/footer.
"""
import io
from docx import Document


def _iter_paragraphs(doc):
    """Duyệt mọi paragraph: thân bài, bảng (kể cả bảng lồng), header, footer."""
    def walk_container(c):
        for p in c.paragraphs:
            yield p
        for t in getattr(c, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk_container(cell)

    yield from walk_container(doc)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer, sec.even_page_header, sec.even_page_footer):
            if part is not None:
                try:
                    yield from walk_container(part)
                except Exception:
                    pass


def read_docx(file_bytes):
    """Trả về (doc, danh_sách_paragraph_object, danh_sách_text)."""
    doc = Document(io.BytesIO(file_bytes))
    paras = [p for p in _iter_paragraphs(doc)]
    texts = [p.text for p in paras]
    return doc, paras, texts


def _replace_in_paragraph(para, start, end, new_text):
    """Thay đoạn [start:end) trong paragraph bằng new_text, giữ định dạng run."""
    runs = para.runs
    if not runs:
        return False
    pos, target = 0, []
    for r in runs:
        rl = len(r.text)
        r_start, r_end = pos, pos + rl
        if r_end > start and r_start < end:
            target.append((r, max(start - r_start, 0), min(end - r_start, rl)))
        pos = r_end
    if not target:
        return False
    # ghi toàn bộ chữ mới vào run đầu tiên, xoá phần tương ứng ở các run sau
    first_run, fs, fe = target[0]
    first_run.text = first_run.text[:fs] + new_text + first_run.text[fe:]
    for r, s, e in target[1:]:
        r.text = r.text[:s] + r.text[e:]
    return True


def apply_to_docx(file_bytes, issues, chosen_ids):
    """Áp dụng sửa lỗi lên file docx gốc, trả về (BytesIO, số_lỗi_đã_sửa)."""
    doc, paras, _ = read_docx(file_bytes)
    chosen = {int(c) for c in chosen_ids}
    by_para = {}
    for iss in issues:
        if iss["id"] in chosen:
            by_para.setdefault(iss["para"], []).append(iss)

    done = 0
    for pi, lst in by_para.items():
        if pi >= len(paras):
            continue
        para = paras[pi]
        # sửa từ cuối lên đầu để không lệch vị trí
        for iss in sorted(lst, key=lambda x: -x["start"]):
            s, e, rep = iss["start"], iss["end"], iss["suggest"]
            if iss["kind"] == "lap_tu":
                rep = ""
                while s > 0 and para.text[s - 1:s] == " ":
                    s -= 1
            elif not rep:
                continue
            if _replace_in_paragraph(para, s, e, rep):
                done += 1

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio, done


def make_docx_from_text(paragraphs, title="Văn bản đã sửa lỗi chính tả"):
    """Tạo file docx mới từ danh sách đoạn (dùng khi nguồn là .txt)."""
    from docx.shared import Pt
    d = Document()
    d.styles["Normal"].font.name = "Times New Roman"
    d.styles["Normal"].font.size = Pt(13)
    for t in paragraphs:
        d.add_paragraph(t)
    bio = io.BytesIO()
    d.save(bio)
    bio.seek(0)
    return bio
