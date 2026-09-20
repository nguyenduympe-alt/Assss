"""Đọc / ghi file Excel điểm - nhận xét."""
import io, re, unicodedata
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


def _norm(s):
    s = str(s).replace("Đ", "D").replace("đ", "d")
    s = unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", s)


ALIAS = {
    "ho_ten": ["hoten", "hovaten", "hovatenhocsinh", "tenhocsinh", "hs", "name", "fullname"],
    "diem": ["diem", "diemtb", "diemtrungbinh", "dtb", "score", "diemso", "tbm", "diemthi"],
    "muc_do": ["mucdo", "mucdodat", "xeploai", "danhgia", "ketqua", "muc", "level", "hoanthanh"],
    "nhan_xet_goc": ["nhanxet", "nhanxetgv", "ghichu", "note", "comment", "nhanxetcuagiaovien"],
    "lop": ["lop", "class", "tenlop"],
    "mon": ["mon", "monhoc", "subject"],
    "ma_hs": ["mahs", "ma", "sobaodanh", "id", "stt2"],
}


def map_columns(cols):
    m = {}
    for c in cols:
        n = _norm(c)
        for key, al in ALIAS.items():
            if n in al or any(n.startswith(a) for a in al):
                if key not in m.values():
                    m[c] = key
                break
    return m


def read_table(file_storage):
    name = (file_storage.filename or "").lower()
    data = file_storage.read()
    bio = io.BytesIO(data)
    if name.endswith(".csv"):
        df = pd.read_csv(bio)
    else:
        df = pd.read_excel(bio)
    df.columns = [str(c).strip() for c in df.columns]
    # bỏ cột Unnamed rỗng
    df = df.loc[:, ~df.columns.str.match(r"^Unnamed")] if len(df.columns) else df
    return df


NORM_MD = {"cht": "CHT", "chuahoanthanh": "CHT", "chuadat": "CHT", "yeu": "CHT", "kem": "CHT",
           "ht": "HT", "hoanthanh": "HT", "dat": "HT", "tb": "HT", "trungbinh": "HT", "kha": "HT",
           "htt": "HTT", "hoanthanhtot": "HTT", "tot": "HTT", "gioi": "HTT", "xuatsac": "HTT"}


def chuan_muc_do(v):
    if v is None or str(v).strip() == "" or str(v).lower() == "nan":
        return None
    return NORM_MD.get(_norm(v))


def export_xlsx(rows, meta=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Nhan xet"
    meta = meta or {}
    hdr = ["STT", "Họ và tên", "Lớp", "Môn", "Điểm", "Mức độ", "Xếp loại", "Nhận xét"]
    ws.append([meta.get("tieu_de", "BẢNG NHẬN XÉT HỌC SINH")])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hdr))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D6EFE6")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, r in enumerate(rows, 1):
        ws.append([i, r.get("ho_ten"), r.get("lop"), r.get("mon"), r.get("diem"),
                   r.get("muc_do"), r.get("xep_loai"), r.get("nhan_xet")])
    widths = [6, 26, 10, 14, 8, 10, 12, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=4):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
        row[1].alignment = Alignment(vertical="top")
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
